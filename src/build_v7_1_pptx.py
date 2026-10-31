import os
import sys
import json
import shutil
import zipfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_v7_1_pptx():
    base_dir = Path(__file__).resolve().parent.parent
    data_path = base_dir / "content" / "deck_data.json"
    deliverables_dir = base_dir / "deliverables"
    assets_dir = base_dir / "src" / "assets"
    deliverables_dir.mkdir(parents=True, exist_ok=True)

    if not data_path.exists():
        raise FileNotFoundError(f"Canonical data file missing: {data_path}")

    with open(data_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    meta = data["meta"]
    slides_data = {s["slide_number"]: s for s in data["slides"]}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ReturnLens V7.1 Refined Palette (Subtle, High-Contrast, No Harsh Shadows)
    C_PAPER = RGBColor(0xFA, 0xF8, 0xF5)          # Warm Paper canvas (#FAF8F5)
    C_PAPER_TINT = RGBColor(0xF3, 0xEF, 0xE9)     # Warm Paper Tint
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)          # Crisp card white
    C_INK = RGBColor(0x11, 0x20, 0x1A)            # Forest Ink (#11201A)
    C_INK_LIGHT = RGBColor(0x28, 0x3C, 0x33)      # Charcoal slate
    C_MUTED = RGBColor(0x56, 0x6B, 0x60)          # Sage Muted
    C_SUBTLE = RGBColor(0x8A, 0x9C, 0x93)         # Hairline gray
    C_BORDER = RGBColor(0xE2, 0xDC, 0xD2)         # Soft Border (#E2DCD2)
    C_BORDER_STRONG = RGBColor(0xCA, 0xC2, 0xB3)  # Strong Divider

    # Accents
    C_TEAL = RGBColor(0x0A, 0x5C, 0x48)           # Pine Teal (#0A5C48)
    C_TEAL_DARK = RGBColor(0x06, 0x40, 0x32)      # Dark Pine
    C_TEAL_LIGHT = RGBColor(0xE8, 0xF5, 0xF1)     # Soft Mint
    C_TEAL_ACCENT = RGBColor(0x2D, 0xD4, 0xBF)    # Mint Accent for dark surfaces

    C_CORAL = RGBColor(0xC8, 0x4B, 0x31)          # Terracotta Coral (#C84B31)
    C_CORAL_DARK = RGBColor(0x9E, 0x2A, 0x2B)     # Deep Coral
    C_CORAL_LIGHT = RGBColor(0xFD, 0xF3, 0xF1)    # Blush Coral Tint
    C_CORAL_BORDER = RGBColor(0xF7, 0xC5, 0xBC)   # Coral Border

    C_AMBER = RGBColor(0xD9, 0x77, 0x06)          # Amber
    C_AMBER_DARK = RGBColor(0x92, 0x40, 0x0E)     # Deep Amber
    C_AMBER_LIGHT = RGBColor(0xFE, 0xF3, 0xC7)    # Amber Cream
    C_GOLD_TEXT = RGBColor(0xFF, 0xD4, 0x85)      # Gold Highlight for dark cards

    # macOS Window Buttons
    C_MAC_RED = RGBColor(0xFF, 0x5F, 0x56)
    C_MAC_YELLOW = RGBColor(0xFF, 0xBD, 0x2E)
    C_MAC_GREEN = RGBColor(0x27, 0xC9, 0x3F)

    # Dark Obsidian Architecture Palette (Slide 4)
    C_DARK_BG = RGBColor(0x0B, 0x16, 0x13)        # Obsidian Canvas (#0B1613)
    C_DARK_CARD = RGBColor(0x13, 0x23, 0x1E)      # Obsidian Card
    C_DARK_BORDER = RGBColor(0x22, 0x3B, 0x33)    # Obsidian Border
    C_DARK_INK = RGBColor(0xF8, 0xF7, 0xF4)       # Pure Warm White text
    C_DARK_MUTED = RGBColor(0x95, 0xAF, 0xA5)     # Sage Muted text

    FONT_HEAD = "Calibri"
    FONT_MONO = "Consolas"

    # Image assets
    img_white_order = assets_dir / "studio_white_headphones.jpg"
    img_black_unboxing = assets_dir / "unboxing_black_headphones.jpg"

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        bg.line.width = 0

    def add_footer(slide, current_num, dark=False):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.46), Inches(7.0), Inches(12.41), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = C_DARK_BORDER if dark else C_BORDER
        line.line.width = 0

        tb = slide.shapes.add_textbox(Inches(0.46), Inches(7.05), Inches(10.0), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "RETURNLENS V7.1 · AGENTX HACKATHON 2026 PS-04 · TEAM ROOKIE (3 MEMBERS)"
        p.font.name = FONT_MONO
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_DARK_MUTED if dark else C_SUBTLE

        tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.37), Inches(0.3))
        tf_num = tb_num.text_frame
        tf_num.word_wrap = False
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"0{current_num} / 06"
        p_num.font.name = FONT_MONO
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = C_DARK_MUTED if dark else C_SUBTLE

    def add_notes(slide, text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = text

    # =========================================================
    # SLIDE 1: PRODUCT HERO & REFINED SELLER DECISION CONSOLE
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, C_PAPER)
    d1 = slides_data[1]

    # Category Tag
    tb = s1.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "AGENTX HACKATHON 2026 · TRACK PS-04 · TEAM ROOKIE (3 MEMBERS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # Left Column: Brand Story, Value Prop & Architectural Principle
    tb_title = s1.shapes.add_textbox(Inches(0.46), Inches(0.85), Inches(4.5), Inches(4.1))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

    p = tf_t.paragraphs[0]
    p.text = d1.get("headline", meta["project_name"])
    p.font.name = FONT_HEAD
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p.space_after = Pt(12)

    p = tf_t.add_paragraph()
    p.text = d1.get("subheadline", meta["tagline"])
    p.font.name = FONT_HEAD
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.space_after = Pt(14)

    p = tf_t.add_paragraph()
    p.text = d1.get("supporting", meta["one_liner"])
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(16)

    # Core Architectural Principle Card (Left Column bottom)
    left_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), Inches(5.12), Inches(4.5), Inches(1.58))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = C_PAPER_TINT
    left_box.line.color.rgb = C_BORDER
    left_box.line.width = Pt(1)

    tb_lb = s1.shapes.add_textbox(Inches(0.62), Inches(5.24), Inches(4.18), Inches(1.34))
    tf_lb = tb_lb.text_frame
    tf_lb.word_wrap = True
    p = tf_lb.paragraphs[0]
    p.text = "CORE ARCHITECTURAL PRINCIPLE"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf_lb.add_paragraph()
    p.text = f"“{meta['core_thesis']}”"
    p.font.name = FONT_HEAD
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_lb.add_paragraph()
    p.text = f"Scope: Return triage + order status · Mock fixtures · Never auto-refunds\nGuardrail: Server cap ≤ min({meta['refund_guardrail']['ceiling']}, remaining, paid) · Hash-bound"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_INK_LIGHT

    # Right Column: Refined macOS Seller Console
    c_x, c_y, c_w = Inches(5.22), Inches(0.85), Inches(7.65)
    c_btn_top = Inches(5.34)
    c_btn_h = Inches(0.38)
    c_btn_bottom = c_btn_top + c_btn_h
    c_pad_bottom = Inches(0.22)
    c_bottom = c_btn_bottom + c_pad_bottom
    c_h = c_bottom - c_y

    c_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, c_w, c_h)
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = C_WHITE
    c_card.line.color.rgb = C_BORDER_STRONG
    c_card.line.width = Pt(1)

    # 1. macOS Window Title Bar with Realistic Traffic-Light Controls
    bar_h = Inches(0.42)
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, c_y, c_w, bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_INK
    bar.line.width = 0

    dot_d = Inches(0.13)
    dot_y = c_y + Inches(0.145)
    dot_colors = [C_MAC_RED, C_MAC_YELLOW, C_MAC_GREEN]
    for i, col in enumerate(dot_colors):
        dot = s1.shapes.add_shape(MSO_SHAPE.OVAL, c_x + Inches(0.20 + i * 0.18), dot_y, dot_d, dot_d)
        dot.fill.solid()
        dot.fill.fore_color.rgb = col
        dot.line.width = 0

    tb_bar = s1.shapes.add_textbox(c_x + Inches(0.80), c_y + Inches(0.06), c_w - Inches(0.95), Inches(0.32))
    tf_b = tb_bar.text_frame
    tf_b.margin_left = tf_b.margin_top = 0
    p = tf_b.paragraphs[0]
    p.text = f"{d1['console_mockup']['title']} — {d1['console_mockup']['case_badge']}       [ MOCK EVIDENCE PACK ]"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFD, 0xF2, 0xD0)

    # 2. Metadata Sub-bar
    sub_y = c_y + bar_h
    sub_h = Inches(0.34)
    subbar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, sub_y, c_w, sub_h)
    subbar.fill.solid()
    subbar.fill.fore_color.rgb = C_PAPER_TINT
    subbar.line.color.rgb = C_BORDER
    subbar.line.width = Pt(1)

    tb_sub = s1.shapes.add_textbox(c_x + Inches(0.2), sub_y + Inches(0.04), c_w - Inches(0.4), Inches(0.28))
    tf_s = tb_sub.text_frame
    tf_s.margin_left = tf_s.margin_top = 0
    p = tf_s.paragraphs[0]
    p.text = f"CASE: {meta['sample_order']['case_id']} · ORDER: {meta['sample_order']['order_id']} · SKU: {meta['sample_order']['sku']}      DELIVERED: 4d ago (Window: Day 4/7 PASS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_INK

    # 3. Side-by-side Visual Comparison
    b_y = sub_y + sub_h + Inches(0.12)
    bx_w = Inches(3.55)
    b_h = Inches(2.55)

    # Left Box - Store Truth
    box_l = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.20), b_y, bx_w, b_h)
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = C_PAPER
    box_l.line.color.rgb = C_BORDER
    box_l.line.width = Pt(1)

    if img_white_order.exists():
        s1.shapes.add_picture(str(img_white_order), c_x + Inches(0.30), b_y + Inches(0.10), bx_w - Inches(0.20), Inches(1.18))

    tb_l = s1.shapes.add_textbox(c_x + Inches(0.30), b_y + Inches(1.32), bx_w - Inches(0.20), Inches(1.15))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "STORE TRUTH · SHOPIFY"
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_MUTED
    p = tf_l.add_paragraph()
    p.text = f"{meta['sample_order']['sku']} ({meta['sample_order']['ordered_color']})"
    p.font.name = FONT_HEAD
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_l.add_paragraph()
    p.text = f"{meta['sample_order']['price']} paid · Studio Spec · Tag present\nDelivered: 4d ago (Shippo #SH-9012)\nReturn Window: Day 4/7 (P-04 PASS)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_MUTED

    # Right Box - Grounded Observation: Black casing, color-consistent with HEAD-BLK-09, SKU not verified
    box_r = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(3.90), b_y, bx_w, b_h)
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = C_CORAL_LIGHT
    box_r.line.color.rgb = C_CORAL_BORDER
    box_r.line.width = Pt(1.5)

    if img_black_unboxing.exists():
        s1.shapes.add_picture(str(img_black_unboxing), c_x + Inches(4.00), b_y + Inches(0.10), bx_w - Inches(0.20), Inches(1.18))

    tb_r = s1.shapes.add_textbox(c_x + Inches(4.00), b_y + Inches(1.32), bx_w - Inches(0.20), Inches(1.15))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "OBSERVED · VIDEO FRAME 00:07"
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p = tf_r.add_paragraph()
    p.text = "Black Casing (color-consistent with HEAD-BLK-09)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK
    p = tf_r.add_paragraph()
    p.text = "Barcode: Unverified · Exact SKU: Not verified\nLighting: Optimal (142/255) · Policy P-03 FAIL\nContradiction flagged → Route to human seller"
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_INK_LIGHT

    # 4. FIX A: Precise P-03 Color Contradiction Notice Strip (Eliminate exact SKU overclaim)
    conf_y = b_y + b_h + Inches(0.10)
    conf_h = Inches(0.44)
    conf = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.20), conf_y, c_w - Inches(0.40), conf_h)
    conf.fill.solid()
    conf.fill.fore_color.rgb = C_CORAL_LIGHT
    conf.line.color.rgb = C_CORAL_BORDER
    conf.line.width = Pt(1)

    tb_c = s1.shapes.add_textbox(c_x + Inches(0.24), conf_y + Inches(0.06), c_w - Inches(0.48), conf_h - Inches(0.10))
    tf_c = tb_c.text_frame
    tf_c.margin_left = Inches(0.04)
    tf_c.margin_right = Inches(0.04)
    tf_c.margin_top = Inches(0.02)
    tf_c.margin_bottom = Inches(0.02)
    p = tf_c.paragraphs[0]
    p.text = "P-03 COLOR CONTRADICTION: Observed black casing conflicts with ordered Matte White specification. Exact SKU unverified."
    p.font.name = FONT_MONO
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK

    # 5. Audit & Ledger Strip
    aud_y = conf_y + conf_h + Inches(0.08)
    aud_h = Inches(0.34)
    aud_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.20), aud_y, c_w - Inches(0.40), aud_h)
    aud_box.fill.solid()
    aud_box.fill.fore_color.rgb = C_PAPER_TINT
    aud_box.line.color.rgb = C_BORDER
    aud_box.line.width = Pt(1)

    tb_aud = s1.shapes.add_textbox(c_x + Inches(0.32), aud_y + Inches(0.05), c_w - Inches(0.64), aud_h - Inches(0.08))
    tf_aud = tb_aud.text_frame
    p = tf_aud.paragraphs[0]
    p.text = f"AUDIT LEDGER: Logged to SQLite · SHA256 evidence hash · Server refund cap: ≤ min({meta['refund_guardrail']['ceiling']}, rem, paid)"
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED

    # 6. Bottom Decision & Action Bar (FIX A: "Review Evidence & Decide")
    tb_act = s1.shapes.add_textbox(c_x + Inches(0.20), c_btn_top + Inches(0.04), Inches(3.2), Inches(0.32))
    tf_a = tb_act.text_frame
    p = tf_a.paragraphs[0]
    p.text = "● SELLER GATE (FAIL-CLOSED)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_AMBER_DARK

    btn_sec = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(3.45), c_btn_top, Inches(1.85), c_btn_h)
    btn_sec.fill.solid()
    btn_sec.fill.fore_color.rgb = C_PAPER_TINT
    btn_sec.line.color.rgb = C_BORDER_STRONG
    btn_sec.line.width = Pt(1)
    p = btn_sec.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Request Box Label"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK

    primary_cta_text = d1['console_mockup'].get('button_text', 'Review Evidence & Decide')
    btn_pri = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(5.40), c_btn_top, Inches(2.05), c_btn_h)
    btn_pri.fill.solid()
    btn_pri.fill.fore_color.rgb = C_TEAL
    btn_pri.line.width = 0
    p = btn_pri.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = primary_cta_text
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    add_footer(s1, 1)
    add_notes(s1, d1["speaker_notes"])

    # =========================================================
    # SLIDE 2: THE PROBLEM (PURPOSE-BUILT INFORMATION FLOW & FIX C)
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, C_PAPER)
    d2 = slides_data[2]

    tb = s2.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "THE PROBLEM · SOLO MERCHANT BOTTLENECK · 4 SOURCES"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s2.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "One return. Four disconnected systems. Zero shared record."
    p.font.name = FONT_HEAD
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = d2.get("supporting_statement", "Solo merchants cross-reference tabs, video footage, and policy rules manually before moving money.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_MUTED

    # Refined Information-Flow Diagram
    dia_x, dia_y, dia_w, dia_h = Inches(0.46), Inches(1.65), Inches(12.41), Inches(3.95)

    src_w = Inches(4.5)
    for i, src in enumerate(d2["sources"]):
        sy = dia_y + Inches(i * 0.98)
        sbox = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dia_x, sy, src_w, Inches(0.84))
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = C_WHITE
        sbox.line.color.rgb = C_BORDER_STRONG
        sbox.line.width = Pt(1)

        tag_bg = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dia_x + Inches(0.12), sy + Inches(0.10), Inches(1.15), Inches(0.24))
        tag_bg.fill.solid()
        tag_bg.fill.fore_color.rgb = C_TEAL_LIGHT
        tag_bg.line.width = 0
        p = tag_bg.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"SOURCE 0{i+1}"
        p.font.name = FONT_MONO
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_DARK

        tb_sname = s2.shapes.add_textbox(dia_x + Inches(1.35), sy + Inches(0.07), src_w - Inches(1.45), Inches(0.26))
        p = tb_sname.text_frame.paragraphs[0]
        p.text = src["name"].upper()
        p.font.name = FONT_HEAD
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_INK

        tb_sdet = s2.shapes.add_textbox(dia_x + Inches(0.15), sy + Inches(0.38), src_w - Inches(0.30), Inches(0.42))
        tf = tb_sdet.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = src["detail"]
        p.font.name = FONT_HEAD
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED

        arr = s2.shapes.add_textbox(dia_x + src_w + Inches(0.10), sy + Inches(0.22), Inches(0.45), Inches(0.40))
        p = arr.text_frame.paragraphs[0]
        p.text = "→"
        p.font.name = FONT_HEAD
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = C_TEAL

    # Right: Prominent Dark Focal Region
    fc_x = dia_x + src_w + Inches(0.65)
    fc_w = dia_w - (src_w + Inches(0.65))
    fc_h = dia_h

    fc = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fc_x, dia_y, fc_w, fc_h)
    fc.fill.solid()
    fc.fill.fore_color.rgb = C_INK
    fc.line.width = 0

    tb_fc = s2.shapes.add_textbox(fc_x + Inches(0.45), dia_y + Inches(0.40), fc_w - Inches(0.90), fc_h - Inches(0.80))
    tf_fc = tb_fc.text_frame
    tf_fc.word_wrap = True
    p = tf_fc.paragraphs[0]
    p.text = "THE CRITICAL INFORMATION GAP"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_TEXT
    p.space_after = Pt(14)

    p = tf_fc.add_paragraph()
    p.text = "“What verifiable evidence justifies moving this money?”"
    p.font.name = FONT_HEAD
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(16)

    p = tf_fc.add_paragraph()
    p.text = "Four disconnected systems. One solo seller making high-stakes financial decisions across browser tabs."
    p.font.name = FONT_HEAD
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(0xDF, 0xEA, 0xE5)
    p.space_after = Pt(12)

    p = tf_fc.add_paragraph()
    p.text = "Without an auditable evidence pack, the merchant either refunds blind (losing COGS and 2-way freight) or disputes unfairly (destroying customer trust and merchant ratings)."
    p.font.name = FONT_HEAD
    p.font.size = Pt(13.5)
    p.font.color.rgb = RGBColor(0xA5, 0xC0, 0xB6)

    # Bottom Row: FIX C — Fraud Metric Clarity (Explicitly Defined & Non-Equated)
    b_y = dia_y + dia_h + Inches(0.18)
    b_w = Inches(6.08)

    bs_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), b_y, b_w, Inches(1.02))
    bs_l.fill.solid()
    bs_l.fill.fore_color.rgb = C_PAPER_TINT
    bs_l.line.color.rgb = C_BORDER
    bs_l.line.width = Pt(1)

    tb_bsl = s2.shapes.add_textbox(Inches(0.62), b_y + Inches(0.10), b_w - Inches(0.32), Inches(0.82))
    tf = tb_bsl.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "INDUSTRY FRAUD BASELINE (DEFINITIONS DIFFER · NOT DIRECTLY COMPARABLE)"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf.add_paragraph()
    p.text = "• 9.0% — Return Fraud (NRF 2025): Strict criminal fraud only (unboxing swaps, empty boxes).\n• 15.14% — Fraud + Policy Abuse (Appriss/Deloitte 2024): Criminal fraud combined with policy abuse."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    bs_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.79), b_y, b_w, Inches(1.02))
    bs_r.fill.solid()
    bs_r.fill.fore_color.rgb = C_PAPER_TINT
    bs_r.line.color.rgb = C_BORDER
    bs_r.line.width = Pt(1)

    tb_bsr = s2.shapes.add_textbox(Inches(6.95), b_y + Inches(0.10), b_w - Inches(0.32), Inches(0.82))
    tf = tb_bsr.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SOLO MERCHANT ASYMMETRIC RISK"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf.add_paragraph()
    p.text = "• Wrong refund = Lost product COGS + 2-way shipping costs (unrecoverable margin loss).\n• Unfair denial = Ruined customer trust, chargeback penalties & degraded seller account health."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    add_footer(s2, 2)
    add_notes(s2, d2["speaker_notes"])

    # =========================================================
    # SLIDE 3: STAR DEMO (HERO EVIDENCE COMPARISON)
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, C_PAPER)
    d3 = slides_data[3]

    tb = s3.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = d3.get("tag", "STAR DEMO · CASE B · CORE BENCHMARK")
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s3.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d3.get("headline", "The order says White. The unboxing video shows Black.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = d3.get("subheadline", "Policy P-03 triggers ESCALATE. No mock label created. No refund draft. No denial without human review.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_MUTED

    # Side-by-side Exhibition Boards
    stg_w = Inches(5.62)
    stg_h = Inches(4.00)
    stg_y = Inches(1.65)

    # Left Board: Store Manifest Truth
    bd_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), stg_y, stg_w, stg_h)
    bd_l.fill.solid()
    bd_l.fill.fore_color.rgb = C_WHITE
    bd_l.line.color.rgb = C_BORDER_STRONG
    bd_l.line.width = Pt(1)

    tb = s3.shapes.add_textbox(Inches(0.70), stg_y + Inches(0.16), stg_w - Inches(0.48), Inches(0.32))
    p = tb.text_frame.paragraphs[0]
    p.text = "ORDER MANIFEST · STORE CATALOG TRUTH (#ORD-7729)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    if img_white_order.exists():
        s3.shapes.add_picture(str(img_white_order), Inches(0.70), stg_y + Inches(0.55), stg_w - Inches(0.48), Inches(2.25))

    tb_spec_l = s3.shapes.add_textbox(Inches(0.70), stg_y + Inches(2.90), stg_w - Inches(0.48), Inches(0.98))
    tf = tb_spec_l.text_frame
    p = tf.paragraphs[0]
    p.text = f"SKU: {meta['sample_order']['sku']} · Spec: {meta['sample_order']['ordered_color']} Casing"
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = f"Price: {meta['sample_order']['price']} paid · Delivered: 4d ago (Shippo #SH-9012)\nReturn Window P-04 (7-Day Policy): PASS (Day 4 of 7)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    # Center Prominent Coral Mismatch Pill
    pill_d = Inches(0.92)
    pill = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.21), stg_y + Inches(1.45), pill_d, pill_d)
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_CORAL
    pill.line.width = 0
    p = pill.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "≠"
    p.font.name = FONT_HEAD
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Right Board: Unboxing Video Grounded Observation
    bd_r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.25), stg_y, stg_w, stg_h)
    bd_r.fill.solid()
    bd_r.fill.fore_color.rgb = RGBColor(0xFF, 0xFD, 0xFC)
    bd_r.line.color.rgb = C_CORAL_BORDER
    bd_r.line.width = Pt(1.5)

    tb = s3.shapes.add_textbox(Inches(7.48), stg_y + Inches(0.16), stg_w - Inches(0.48), Inches(0.32))
    p = tb.text_frame.paragraphs[0]
    p.text = "MOCK EVIDENCE PACK · VIDEO CITATION (FRAME 00:07)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL

    if img_black_unboxing.exists():
        s3.shapes.add_picture(str(img_black_unboxing), Inches(7.48), stg_y + Inches(0.55), stg_w - Inches(0.48), Inches(2.25))

    tb_spec_r = s3.shapes.add_textbox(Inches(7.48), stg_y + Inches(2.90), stg_w - Inches(0.48), Inches(0.98))
    tf = tb_spec_r.text_frame
    p = tf.paragraphs[0]
    p.text = "Observed: Black Casing (color-consistent with HEAD-BLK-09)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK
    p = tf.add_paragraph()
    p.text = "Exact SKU: Not verified (Barcode unverified) · Policy P-03 FAIL\nLighting: Optimal (142/255) · REC 00:07 / 00:24 (Mock Keyframe)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_CORAL_DARK

    # Bottom Decision Strip: Prominent Result Banner with 3 Clean Action Pills
    dec_y = stg_y + stg_h + Inches(0.18)
    dec_h = Inches(1.02)
    dec = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), dec_y, Inches(12.41), dec_h)
    dec.fill.solid()
    dec.fill.fore_color.rgb = C_WHITE
    dec.line.color.rgb = C_BORDER_STRONG
    dec.line.width = Pt(1)

    tb_ot = s3.shapes.add_textbox(Inches(0.65), dec_y + Inches(0.08), Inches(6.5), Inches(0.30))
    p = tb_ot.text_frame.paragraphs[0]
    p.text = "POLICY P-03 RESULT: ESCALATE TO SELLER GATE (FAIL-CLOSED)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK

    tb_fg = s3.shapes.add_textbox(Inches(6.80), dec_y + Inches(0.08), Inches(5.8), Inches(0.30))
    p = tb_fg.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = f"Guardrail: 0 < amt ≤ min({meta['refund_guardrail']['ceiling']}, rem_ref, paid) · Never auto-refunds"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # 3 Preserved Action Interactive Buttons
    actions = [
        "1. Accept as warehouse mis-shipment",
        "2. Request photo of outside-box shipping label",
        "3. Hold return for physical parcel inspection"
    ]
    pill_w = Inches(3.92)
    for idx, act in enumerate(actions):
        ax = Inches(0.65 + idx * 4.05)
        abtn = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ax, dec_y + Inches(0.45), pill_w, Inches(0.44))
        abtn.fill.solid()
        abtn.fill.fore_color.rgb = C_PAPER_TINT
        abtn.line.color.rgb = C_BORDER_STRONG
        abtn.line.width = Pt(1)
        p = abtn.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = act
        p.font.name = FONT_HEAD
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_INK

    add_footer(s3, 3)
    add_notes(s3, d3["speaker_notes"])

    # =========================================================
    # SLIDE 4: SYSTEM ARCHITECTURE (OBSIDIAN DARK THEME & FIX B)
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, C_DARK_BG)
    d4 = slides_data[4]

    tb = s4.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = d4.get("tag", "SYSTEM ARCHITECTURE · FAIL-CLOSED GOVERNANCE")
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT

    tb = s4.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d4.get("headline", meta["core_thesis"])
    p.font.name = FONT_HEAD
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p = tf.add_paragraph()
    p.text = "Exactly 3 base model calls per triage case. Weak evidence escalates. Never releases automated refunds."
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_DARK_MUTED

    # 6-Stage Visual Architecture Pipeline with Distinct Authority Types & Directional Flow
    pip_y = Inches(1.65)
    pip_h = Inches(1.10)
    pip = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), pip_y, Inches(12.41), pip_h)
    pip.fill.solid()
    pip.fill.fore_color.rgb = C_DARK_CARD
    pip.line.color.rgb = C_DARK_BORDER
    pip.line.width = Pt(1)

    steps = [
        ("1", "Planner", "Model Call 1", "Proposes typed JSON actions"),
        ("2", "Tools & Vision", "Model Call 2", "FFmpeg keyframes + vision"),
        ("3", "Policy Engine", "Deterministic Code", "Evaluates rules P-01–P-08"),
        ("4", "Critic", "Model Call 3", "Verifies cited timestamps"),
        ("5", "Seller Gate", "Human Merchant", "Fail-closed 1-click auth"),
        ("6", "Audit Ledger", "SQLite / JSON", "Append-only signed audit")
    ]
    for idx, (num, name, role_badge, desc) in enumerate(steps):
        step_x = Inches(0.65 + idx * 2.05)
        tb_step = s4.shapes.add_textbox(step_x, pip_y + Inches(0.10), Inches(1.90), Inches(0.90))
        tf = tb_step.text_frame
        p = tf.paragraphs[0]
        p.text = f"{num}. {name}"
        p.font.name = FONT_MONO
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_ACCENT if "Model" in role_badge else C_DARK_INK
        p = tf.add_paragraph()
        p.text = role_badge
        p.font.name = FONT_MONO
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = C_GOLD_TEXT if "Human" in role_badge else (C_TEAL_ACCENT if "Model" in role_badge else C_DARK_MUTED)
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT_HEAD
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK_MUTED

    # 3 Floating Architectural Pillars
    g_y = Inches(2.92)
    g_w = Inches(3.95)
    g_h = Inches(2.72)

    # Pillar 1: Call Budget & Zero Payment Access
    gc1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), g_y, g_w, g_h)
    gc1.fill.solid()
    gc1.fill.fore_color.rgb = C_DARK_CARD
    gc1.line.color.rgb = C_DARK_BORDER
    gc1.line.width = Pt(1)

    tb = s4.shapes.add_textbox(Inches(0.65), g_y + Inches(0.18), g_w - Inches(0.40), g_h - Inches(0.36))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CALL BUDGET & MODEL RESTRICTION"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Exactly 3 Model Calls"
    p.font.name = FONT_HEAD
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "1. Planner proposes typed tool actions.\n2. Vision model inspects server-extracted keyframes.\n3. Critic evaluates citations against policy.\n\nMODELS NEVER ACCESS PAYMENT APIS DIRECTLY."
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.color.rgb = C_DARK_MUTED

    # Pillar 2: FIX B — Case C Stateful Resumption Diagram & P-06 Failure
    gc2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.69), g_y, g_w, g_h)
    gc2.fill.solid()
    gc2.fill.fore_color.rgb = RGBColor(0x14, 0x28, 0x22)
    gc2.line.color.rgb = C_TEAL
    gc2.line.width = Pt(1.5)

    tb = s4.shapes.add_textbox(Inches(4.88), g_y + Inches(0.18), g_w - Inches(0.40), g_h - Inches(0.36))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "STATEFUL RESUMPTION (CASE C)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Clarification & Re-Film Loop"
    p.font.name = FONT_HEAD
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "Low-light video (Policy P-06 FAIL, brightness score: 28/255) triggers ask_for_evidence().\n\nState saves to WAITING_FOR_EVIDENCE → Resumes on upload.\n\nTERMINATION RULE: After 1 retry, escalates to seller with both attempts logged (never auto-denies)."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11.5)
    p.font.color.rgb = RGBColor(0xD2, 0xE5, 0xDC)

    # Pillar 3: Secondary Fast Path (Tiny PS-04 Order Route)
    gc3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.92), g_y, g_w, g_h)
    gc3.fill.solid()
    gc3.fill.fore_color.rgb = C_DARK_CARD
    gc3.line.color.rgb = C_DARK_BORDER
    gc3.line.width = Pt(1)

    tb = s4.shapes.add_textbox(Inches(9.11), g_y + Inches(0.18), g_w - Inches(0.40), g_h - Inches(0.36))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SECONDARY PS-04 ROUTE"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Tiny PS-04 Order Route"
    p.font.name = FONT_HEAD
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "For routine \"Where is my order?\" inquiries, the agent bypasses vision models entirely:\n\nStore API → Courier API → Status Reply\n\nZero vision token waste · Instant response."
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.color.rgb = C_DARK_MUTED

    # Bottom Technical Banner
    ban_y = g_y + g_h + Inches(0.16)
    ban = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), ban_y, Inches(12.41), Inches(0.80))
    ban.fill.solid()
    ban.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x1D)
    ban.line.color.rgb = C_DARK_BORDER
    ban.line.width = Pt(1)

    tb_ban = s4.shapes.add_textbox(Inches(0.65), ban_y + Inches(0.10), Inches(12.0), Inches(0.60))
    tf = tb_ban.text_frame
    p = tf.paragraphs[0]
    p.text = "FAIL-CLOSED FINANCIAL GATE: requires_seller_approval: true · Approval binds case + evidence hash + policy version."
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_TEXT
    p = tf.add_paragraph()
    p.text = f"Server Cap: 0 < amount <= min({meta['refund_guardrail']['ceiling']}, remaining_refundable, order_paid) · Revalidation check before draft release."
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_DARK_MUTED

    add_footer(s4, 4, dark=True)
    add_notes(s4, d4["speaker_notes"])

    # =========================================================
    # SLIDE 5: 4 EVALUATION SCENARIOS (ASYMMETRIC MATRIX & FIX A/B)
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, C_PAPER)
    d5 = slides_data[5]

    tb = s5.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = d5.get("tag", "EVALUATION SUITE · 4 CONTROLLED SCENARIOS")
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s5.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d5.get("headline", "Different evidence. Controlled, predictable action.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = d5.get("subheadline", "Every case produces an auditable mock evidence pack. No automatic policy approvals.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_MUTED

    # Asymmetric Layout:
    b_x, b_y, b_w, b_h = Inches(0.46), Inches(1.65), Inches(5.40), Inches(5.12)
    bc = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_x, b_y, b_w, b_h)
    bc.fill.solid()
    bc.fill.fore_color.rgb = C_WHITE
    bc.line.color.rgb = C_CORAL_BORDER
    bc.line.width = Pt(2)

    # 1. Title Block
    tb_title = s5.shapes.add_textbox(b_x + Inches(0.24), b_y + Inches(0.16), b_w - Inches(0.48), Inches(0.76))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p = tf_t.paragraphs[0]
    p.text = "CORE BENCHMARK SCENARIO · P-03 CONTRADICTION"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p = tf_t.add_paragraph()
    p.text = "Case B: Color Contradiction"
    p.font.name = FONT_HEAD
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_t.add_paragraph()
    p.text = "The order says White. Unboxing video frame 00:07 shows Black."
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.color.rgb = C_MUTED

    # 2. Keyframe Evidence Citation Box with Mini Thumbnail
    ev_y = b_y + Inches(0.98)
    ev_h = Inches(1.52)
    ev_w = b_w - Inches(0.48)
    ev_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_x + Inches(0.24), ev_y, ev_w, ev_h)
    ev_box.fill.solid()
    ev_box.fill.fore_color.rgb = C_CORAL_LIGHT
    ev_box.line.color.rgb = C_CORAL_BORDER
    ev_box.line.width = Pt(1)

    if img_black_unboxing.exists():
        s5.shapes.add_picture(str(img_black_unboxing), b_x + Inches(3.50), ev_y + Inches(0.12), Inches(1.55), Inches(1.28))

    tb_ev = s5.shapes.add_textbox(b_x + Inches(0.36), ev_y + Inches(0.08), Inches(3.10), ev_h - Inches(0.14))
    tf_ev = tb_ev.text_frame
    tf_ev.word_wrap = True
    tf_ev.margin_left = tf_ev.margin_top = tf_ev.margin_right = tf_ev.margin_bottom = 0
    p = tf_ev.paragraphs[0]
    p.text = "MOCK EVIDENCE PACK · FRAME 00:07"
    p.font.name = FONT_MONO
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p = tf_ev.add_paragraph()
    p.text = f"Manifest: {meta['sample_order']['sku']} White ({meta['sample_order']['price']})"
    p.font.name = FONT_HEAD
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_INK
    p = tf_ev.add_paragraph()
    p.text = "Observed: Black casing (color-consistent)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK
    p = tf_ev.add_paragraph()
    p.text = "Barcode: Unverified · Exact SKU: Not verified"
    p.font.name = FONT_MONO
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_MUTED
    p = tf_ev.add_paragraph()
    p.text = "P-03 FAIL → No label, no refund draft"
    p.font.name = FONT_HEAD
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK

    # 3. Policy Action Strip
    dec_y = ev_y + ev_h + Inches(0.08)
    dec_h = Inches(0.42)
    dec_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_x + Inches(0.24), dec_y, ev_w, dec_h)
    dec_box.fill.solid()
    dec_box.fill.fore_color.rgb = C_INK
    dec_box.line.width = 0

    tb_dec = s5.shapes.add_textbox(b_x + Inches(0.36), dec_y + Inches(0.08), ev_w - Inches(0.24), dec_h - Inches(0.14))
    tf_dec = tb_dec.text_frame
    tf_dec.margin_left = tf_dec.margin_top = 0
    p = tf_dec.paragraphs[0]
    p.text = "POLICY ACTION: ESCALATE TO SELLER GATE (FAIL-CLOSED)"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_TEXT

    # 4. Preserved Seller Options Box
    opt_y = dec_y + dec_h + Inches(0.08)
    opt_h = Inches(1.85)
    opt_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_x + Inches(0.24), opt_y, ev_w, opt_h)
    opt_box.fill.solid()
    opt_box.fill.fore_color.rgb = C_PAPER_TINT
    opt_box.line.color.rgb = C_BORDER
    opt_box.line.width = Pt(1)

    tb_opt = s5.shapes.add_textbox(b_x + Inches(0.36), opt_y + Inches(0.10), ev_w - Inches(0.24), opt_h - Inches(0.16))
    tf_opt = tb_opt.text_frame
    tf_opt.word_wrap = True
    tf_opt.margin_left = tf_opt.margin_top = tf_opt.margin_right = tf_opt.margin_bottom = 0
    p = tf_opt.paragraphs[0]
    p.text = "PRESERVED SELLER OPTIONS IN DECISION CONSOLE:"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf_opt.add_paragraph()
    p.text = "① Accept warehouse mis-shipment"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_opt.add_paragraph()
    p.text = "② Request photo of outside-box shipping label"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_opt.add_paragraph()
    p.text = "③ Hold return for physical parcel inspection"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_opt.add_paragraph()
    p.text = f"Financial Guardrail: 0 < amount ≤ min({meta['refund_guardrail']['ceiling']}, rem, paid) · Binds evidence hash"
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED

    # Right: Compact Stack of Cases A, C, D (Editorial Mini-Stories)
    rw_x = Inches(6.08)
    rw_w = Inches(6.79)
    rw_h = Inches(1.58)

    cases_right = [
        ("Case A: Valid Return (ORD-1042)", "PREPARE SELLER REVIEW", C_TEAL_DARK, C_TEAL_LIGHT,
         "Evidence: ORD-1042 · ₹1,499 · SKU matches catalog · Packaging tag present · Within 7-day return window.",
         "Outcome: Mock shipping label created · Refund draft placed in pending_return_receipt (not executed)."),
        ("Case C: Low-Light Video (Re-Film Loop)", "ASK & RESUME", C_AMBER_DARK, C_AMBER_LIGHT,
         "Evidence: 4-second dark unboxing video · Brightness score: 28/255 · Only 2 usable keyframes extracted.",
         "Outcome: Agent issues re-film checklist under Policy P-06 · State saved to WAITING_FOR_EVIDENCE · Resumes on upload (after 1 retry, escalates with both attempts logged; never auto-denies)."),
        ("Case D: Continuous Unboxing Sequence", "INSPECT SEQUENCE", C_INK, C_PAPER_TINT,
         "Evidence: Single continuous take from opening sealed parcel to displaying label in one sequence.",
         "Outcome: Inspects sequence continuity and flags abrupt transitions. Does not claim proof of identical item.")
    ]

    for idx, (title, badge, b_col, bg_col, ev, out) in enumerate(cases_right):
        cy = b_y + Inches(idx * 1.76)
        cr = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rw_x, cy, rw_w, rw_h)
        cr.fill.solid()
        cr.fill.fore_color.rgb = C_WHITE
        cr.line.color.rgb = C_BORDER
        cr.line.width = Pt(1)

        tb = s5.shapes.add_textbox(rw_x + Inches(0.25), cy + Inches(0.12), rw_w - Inches(0.50), rw_h - Inches(0.20))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}   [{badge}]"
        p.font.name = FONT_HEAD
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = b_col
        p = tf.add_paragraph()
        p.text = ev
        p.font.name = FONT_HEAD
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED
        p = tf.add_paragraph()
        p.text = out
        p.font.name = FONT_MONO
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_DARK

    add_footer(s5, 5)
    add_notes(s5, d5["speaker_notes"])

    # =========================================================
    # SLIDE 6: 24-HOUR PROTOTYPE EXECUTION PLAN
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, C_PAPER)
    d6 = slides_data[6]

    tb = s6.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "24-HOUR PROTOTYPE PLAN · TEAM ROOKIE (3 MEMBERS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s6.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d6.get("headline", "A credible, buildable prototype in 24 hours.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = d6.get("subheadline", "Strict scope discipline, mock data fixtures, and single-threaded ownership ensure on-time delivery.")
    p.font.name = FONT_HEAD
    p.font.size = Pt(15.5)
    p.font.color.rgb = C_MUTED

    # High-Quality Horizontal Roadmap Timeline (00-24h)
    t_y = Inches(1.65)
    t_h = Inches(1.22)
    t_bar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), t_y, Inches(12.41), t_h)
    t_bar.fill.solid()
    t_bar.fill.fore_color.rgb = C_WHITE
    t_bar.line.color.rgb = C_BORDER_STRONG
    t_bar.line.width = Pt(1)

    t_slots = [
        ("00 – 06h", "Foundation", "Mock Store/Shippo APIs · LangGraph state schema · Policy P-01 to P-08"),
        ("06 – 12h", "Evidence Pipeline", "FFmpeg ~1fps extractor · Vision prompt schemas · Seller Console UI mock"),
        ("12 – 18h", "Integration", "Case C pause/resume loop · 3 core cases (A–C) + Case D sequence test · Critic review"),
        ("18 – 24h", "Validation & Safety", "Fail-closed authorization tests · Boundary QA · Video demo recording")
    ]
    for idx, (hrs, phase, desc) in enumerate(t_slots):
        tx = Inches(0.65 + idx * 3.08)

        mbadge = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, t_y + Inches(0.12), Inches(1.65), Inches(0.24))
        mbadge.fill.solid()
        mbadge.fill.fore_color.rgb = C_TEAL_LIGHT
        mbadge.line.width = 0
        p = mbadge.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"{hrs} · PHASE 0{idx+1}"
        p.font.name = FONT_MONO
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_DARK

        tb = s6.shapes.add_textbox(tx, t_y + Inches(0.38), Inches(2.90), Inches(0.76))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = phase
        p.font.name = FONT_HEAD
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_INK
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT_HEAD
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_MUTED

        if idx < 3:
            div = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65 + (idx + 1) * 3.08 - 0.15), t_y + Inches(0.18), Inches(0.015), Inches(0.86))
            div.fill.solid()
            div.fill.fore_color.rgb = C_BORDER
            div.line.width = 0

    # 3 Enriched Person Cards: Exactly 4 Concrete Deliverables Each & Clear Tech Stacks
    p_y = Inches(3.05)
    p_w = Inches(3.95)
    p_h = Inches(2.45)

    people = [
        ("PERSON A · BACKEND, POLICY & AUDIT", "State Machine & Policy Engine",
         "Stack: LangGraph · Pydantic · SQLite · SHA256",
         [
             "• Policy P-01 to P-08: Deterministic code rules (window, SKU, tag, refund cap)",
             "• Critic Agent: Adversarial check against cited keyframe timestamps",
             "• Audit Ledger: Append-only JSON log & cryptographic approval binding",
             "• Financial Cap: Server calculation min(₹4,999, remaining, paid)"
         ]),
        ("PERSON B · VISION & CONSOLE", "Video Pipeline & Seller UI",
         "Stack: FFmpeg · Hosted vision model (unpinned) · HTML/CSS/SVG",
         [
             "• Keyframe Pipeline: Uniform ~1fps sampling & brightness scoring (Case C sample: 28/255)",
             "• Structured Vision: Observation schema with timestamp citation metadata",
             "• Seller Console: 1-click decision console with mock evidence pack",
             "• Grounding Guard: Enforces unverified barcode vs color distinction"
         ]),
        ("PERSON C · INTEGRATION & QA", "Scenarios, Lifecycle & QA",
         "Stack: Pytest · Mock Fixtures · Automated Smoke QA",
         [
             "• Benchmark Cases: Staged data for Cases A–C + Case D sequence test",
             "• Case C Lifecycle: Re-film loop in WAITING_FOR_EVIDENCE (1 retry)",
             "• Fail-Closed Testing: Verifies zero automated refunds & approval gates",
             "• Automated QA: 10-test smoke suite, end-to-end event bus & demo recording"
         ])
    ]

    for idx, (role, name, stack, deliverables) in enumerate(people):
        px = Inches(0.46 + idx * 4.23)
        pc = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, p_y, p_w, p_h)
        pc.fill.solid()
        pc.fill.fore_color.rgb = C_PAPER
        pc.line.color.rgb = C_BORDER
        pc.line.width = Pt(1)

        tb = s6.shapes.add_textbox(px + Inches(0.18), p_y + Inches(0.14), p_w - Inches(0.36), p_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = role
        p.font.name = FONT_MONO
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_TEAL
        p = tf.add_paragraph()
        p.text = name
        p.font.name = FONT_HEAD
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_INK
        p = tf.add_paragraph()
        p.text = stack
        p.font.name = FONT_MONO
        p.font.size = Pt(9)
        p.font.color.rgb = C_MUTED
        p.space_after = Pt(4)
        for item in deliverables:
            p = tf.add_paragraph()
            p.text = item
            p.font.name = FONT_HEAD
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_INK_LIGHT
            p.space_after = Pt(2)

    # Bottom Scope Discipline Band
    sc_y = Inches(5.55)
    sc_h = Inches(1.25)
    sc = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), sc_y, Inches(12.41), sc_h)
    sc.fill.solid()
    sc.fill.fore_color.rgb = C_INK
    sc.line.width = 0

    tb_sc1 = s6.shapes.add_textbox(Inches(0.65), sc_y + Inches(0.12), Inches(4.50), Inches(1.00))
    tf1 = tb_sc1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "MUST DELIVER (CORE PROTOTYPE)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_TEXT
    p = tf1.add_paragraph()
    p.text = "Mock APIs · Policy engine · Cases A–C + Case D sequence test · Approval binding · Seller gate · Tiny order route"
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(0xDF, 0xEA, 0xE5)

    tb_sc2 = s6.shapes.add_textbox(Inches(5.40), sc_y + Inches(0.12), Inches(4.20), Inches(1.00))
    tf2 = tb_sc2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "CUT FIRST IF BEHIND"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0x8C, 0x69)
    p = tf2.add_paragraph()
    p.text = "Advanced video forensics (Case D stretch) · WhatsApp bot · Exchanges · Multi-seller auth · Audio analysis · Public cloud deployment"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11.5)
    p.font.color.rgb = RGBColor(0xDF, 0xEA, 0xE5)

    tb_sc3 = s6.shapes.add_textbox(Inches(9.80), sc_y + Inches(0.12), Inches(2.90), Inches(1.00))
    tf3 = tb_sc3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "FINANCIAL CEILING FORMULA"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf3.add_paragraph()
    p.text = f"0 < amt <= min({meta['refund_guardrail']['ceiling']}, rem_ref, paid)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_WHITE

    add_footer(s6, 6)
    add_notes(s6, d6["speaker_notes"])

    pptx_deliv = deliverables_dir / "ReturnLens_v7_1.pptx"
    pptx_root = base_dir / "ReturnLens_v7_1.pptx"
    temp_raw = deliverables_dir / "temp_raw.pptx"
    prs.save(str(temp_raw))

    # Write deterministic archive with normalized timestamps
    fixed_time = (2026, 10, 31, 12, 0, 0)
    with zipfile.ZipFile(temp_raw, 'r') as zin:
        with zipfile.ZipFile(pptx_deliv, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                new_info = zipfile.ZipInfo(item.filename, date_time=fixed_time)
                new_info.compress_type = item.compress_type
                new_info.external_attr = 0o644 << 16
                zout.writestr(new_info, data)

    if temp_raw.exists():
        temp_raw.unlink()

    shutil.copy2(pptx_deliv, pptx_root)
    print(f"Generated ReturnLens_v7_1.pptx successfully: {pptx_deliv} ({os.path.getsize(pptx_deliv)} bytes)")

if __name__ == "__main__":
    build_v7_1_pptx()
