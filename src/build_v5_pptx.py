import os
import sys
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_v5_pptx():
    base_dir = Path(__file__).resolve().parent.parent
    data_path = base_dir / "content" / "deck_data.json"
    deliverables_dir = base_dir / "deliverables"
    deliverables_dir.mkdir(parents=True, exist_ok=True)

    if not data_path.exists():
        raise FileNotFoundError(f"Canonical data file missing: {data_path}")

    with open(data_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    meta = data["meta"]
    slides_data = {s["slide_number"]: s for s in data["slides"]}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # V5 Color Palette
    C_PAPER = RGBColor(0xFA, 0xF8, 0xF5)
    C_PAPER_TINT = RGBColor(0xF2, 0xEF, 0xEB)
    C_INK = RGBColor(0x14, 0x24, 0x1E)
    C_INK_LIGHT = RGBColor(0x24, 0x38, 0x30)
    C_MUTED = RGBColor(0x5C, 0x70, 0x66)
    C_SUBTLE = RGBColor(0x8A, 0x9C, 0x93)
    C_BORDER = RGBColor(0xE3, 0xDD, 0xD4)
    C_BORDER_STRONG = RGBColor(0xCD, 0xC5, 0xB8)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    C_TEAL = RGBColor(0x00, 0x6B, 0x54)
    C_TEAL_DARK = RGBColor(0x00, 0x4D, 0x3C)
    C_TEAL_LIGHT = RGBColor(0xE6, 0xF4, 0xF0)
    C_TEAL_ACCENT = RGBColor(0x00, 0xD1, 0xA7)

    C_CORAL = RGBColor(0xC0, 0x39, 0x2B)
    C_CORAL_DARK = RGBColor(0x96, 0x2D, 0x22)
    C_CORAL_LIGHT = RGBColor(0xFD, 0xF2, 0xF2)
    C_CORAL_BORDER = RGBColor(0xF5, 0xB7, 0xB1)

    C_AMBER = RGBColor(0xD9, 0x77, 0x06)
    C_AMBER_DARK = RGBColor(0x92, 0x40, 0x0E)
    C_AMBER_LIGHT = RGBColor(0xFE, 0xF3, 0xC7)

    C_DARK_BG = RGBColor(0x0E, 0x18, 0x15)
    C_DARK_CARD = RGBColor(0x15, 0x24, 0x20)
    C_DARK_BORDER = RGBColor(0x23, 0x39, 0x32)
    C_DARK_INK = RGBColor(0xF7, 0xF5, 0xF0)
    C_DARK_MUTED = RGBColor(0x8E, 0xAB, 0xA0)

    FONT_HEAD = "Inter"
    FONT_MONO = "IBM Plex Mono"

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        bg.line.width = 0

    def add_footer(slide, current_num, dark=False):
        # Footer dividing line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.46), Inches(7.0), Inches(12.41), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = C_DARK_BORDER if dark else C_BORDER
        line.line.width = 0

        # Footer text
        tb = slide.shapes.add_textbox(Inches(0.46), Inches(7.05), Inches(10.0), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "RETURNLENS · AGENTX HACKATHON 2026 PS-04 · TEAM ROOKIE (3 MEMBERS)"
        p.font.name = FONT_MONO
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_DARK_MUTED if dark else C_SUBTLE

        tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.37), Inches(0.3))
        tf_num = tb_num.text_frame
        tf_num.word_wrap = False
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"0{current_num} / 06"
        p_num.font.name = FONT_MONO
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = C_DARK_MUTED if dark else C_SUBTLE

    def add_notes(slide, text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = text

    # =========================================================
    # SLIDE 1: PRODUCT HERO & CONSOLE
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, C_PAPER)
    d1 = slides_data[1]

    # Tag
    tb = s1.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "AGENTX HACKATHON 2026 · PS-04 TRACK · TEAM ROOKIE (3 MEMBERS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # Left Column: Story & Meta
    tb_title = s1.shapes.add_textbox(Inches(0.46), Inches(0.85), Inches(4.5), Inches(5.8))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

    p = tf_t.paragraphs[0]
    p.text = "ReturnLens"
    p.font.name = FONT_HEAD
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p.space_after = Pt(14)

    p = tf_t.add_paragraph()
    p.text = "Every return decision, backed by verifiable evidence."
    p.font.name = FONT_HEAD
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.space_after = Pt(16)

    p = tf_t.add_paragraph()
    p.text = "For solo D2C merchants managing orders alone: an autonomous agent that unifies customer video, store facts, and return policy into a 1-click decision console."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_MUTED
    p.space_after = Pt(20)

    p = tf_t.add_paragraph()
    p.text = "Core Thesis: Models observe. Code verifies. Seller authorizes.\nScope: Return triage + order status · Mock fixtures · Never auto-refunds"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    # Right Column: Hero Console Window
    c_x, c_y, c_w, c_h = Inches(5.2), Inches(0.85), Inches(7.67), Inches(5.85)
    c_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, c_w, c_h)
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = C_WHITE
    c_card.line.color.rgb = C_BORDER_STRONG
    c_card.line.width = Pt(1)

    # Console Title Bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, c_y, c_w, Inches(0.52))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_INK
    bar.line.width = 0

    tb_bar = s1.shapes.add_textbox(c_x + Inches(0.2), c_y + Inches(0.08), c_w - Inches(0.4), Inches(0.36))
    tf_b = tb_bar.text_frame
    tf_b.margin_left = tf_b.margin_top = 0
    p = tf_b.paragraphs[0]
    p.text = "RETURNLENS SELLER CONSOLE — CASE #RT-8821               [ MOCK EVIDENCE PACK ]"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFD, 0xF2, 0xD0)

    # Subbar
    subbar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x, c_y + Inches(0.52), c_w, Inches(0.42))
    subbar.fill.solid()
    subbar.fill.fore_color.rgb = C_PAPER_TINT
    subbar.line.color.rgb = C_BORDER
    subbar.line.width = Pt(1)

    tb_sub = s1.shapes.add_textbox(c_x + Inches(0.2), c_y + Inches(0.56), c_w - Inches(0.4), Inches(0.32))
    tf_s = tb_sub.text_frame
    tf_s.margin_left = tf_s.margin_top = 0
    p = tf_s.paragraphs[0]
    p.text = "CASE: #RT-8821 · ORDER: #ORD-7729 · SKU: HEAD-WHT-02      DELIVERED: 4d ago (Window: Day 4/7 PASS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.color.rgb = C_INK

    # Visual Comparison Left (White) & Right (Black)
    bx_w = Inches(3.55)
    b_y = c_y + Inches(1.15)
    b_h = Inches(2.2)

    # Left Box
    box_l = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.2), b_y, bx_w, b_h)
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = C_PAPER
    box_l.line.color.rgb = C_BORDER
    box_l.line.width = Pt(1)

    tb_l = s1.shapes.add_textbox(c_x + Inches(0.35), b_y + Inches(0.2), bx_w - Inches(0.3), b_h - Inches(0.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "STORE TRUTH · SHOPIFY"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_MUTED
    p = tf_l.add_paragraph()
    p.text = "HEAD-WHT-02 (White)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf_l.add_paragraph()
    p.text = "₹1,499 · Studio Catalog Spec\nTag Required: Present in manifest\nStatus: Verified against order"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_MUTED

    # Right Box
    box_r = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(3.92), b_y, bx_w, b_h)
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = C_CORAL_LIGHT
    box_r.line.color.rgb = C_CORAL_BORDER
    box_r.line.width = Pt(1.5)

    tb_r = s1.shapes.add_textbox(c_x + Inches(4.07), b_y + Inches(0.2), bx_w - Inches(0.3), b_h - Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "OBSERVED · VIDEO FRAME 00:07"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p = tf_r.add_paragraph()
    p.text = "HEAD-BLK-09 (Midnight Black)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK
    p = tf_r.add_paragraph()
    p.text = "Observed: Black casing\nExact Barcode: Unverified in frame\nLighting: Optimal (142/255)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    # Conflict Strip
    conf_y = b_y + b_h + Inches(0.2)
    conf = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.2), conf_y, c_w - Inches(0.4), Inches(0.72))
    conf.fill.solid()
    conf.fill.fore_color.rgb = C_CORAL_LIGHT
    conf.line.color.rgb = C_CORAL_BORDER
    conf.line.width = Pt(1)

    tb_c = s1.shapes.add_textbox(c_x + Inches(0.35), conf_y + Inches(0.12), c_w - Inches(0.7), Inches(0.48))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "P-03 CONTRADICTION: Item in parcel does not match ordered SKU · Citation: Keyframe 00:07"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK

    # Bottom Actions Bar
    act_y = conf_y + Inches(0.9)
    tb_act = s1.shapes.add_textbox(c_x + Inches(0.2), act_y, Inches(4.5), Inches(0.4))
    tf_a = tb_act.text_frame
    p = tf_a.paragraphs[0]
    p.text = "● SELLER AUTHORIZATION REQUIRED (FAIL-CLOSED)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_AMBER_DARK

    btn = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(4.9), act_y - Inches(0.05), Inches(2.55), Inches(0.48))
    btn.fill.solid()
    btn.fill.fore_color.rgb = C_TEAL
    btn.line.width = 0
    tf_btn = btn.text_frame
    p = tf_btn.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Review Evidence & Authorize"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    add_footer(s1, 1)
    add_notes(s1, d1["speaker_notes"])

    # =========================================================
    # SLIDE 2: THE PROBLEM (HORIZONTAL CONVERGENCE)
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, C_PAPER)
    d2 = slides_data[2]

    # Tag & Headline
    tb = s2.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "THE PROBLEM · SOLO MERCHANT BOTTLENECK"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s2.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "One return. Four disconnected systems. Zero shared record."
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = "Solo merchants cross-reference tabs, video footage, and policy rules manually before moving money."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_MUTED

    # Diagram Outer Box
    dia_x, dia_y, dia_w, dia_h = Inches(0.46), Inches(1.65), Inches(12.41), Inches(3.95)
    dia = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dia_x, dia_y, dia_w, dia_h)
    dia.fill.solid()
    dia.fill.fore_color.rgb = C_WHITE
    dia.line.color.rgb = C_BORDER
    dia.line.width = Pt(1)

    # Left: 4 Sources
    src_w = Inches(4.3)
    for i, src in enumerate(d2["sources"]):
        sy = dia_y + Inches(0.25 + i * 0.88)
        sbox = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dia_x + Inches(0.25), sy, src_w, Inches(0.72))
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = C_PAPER
        sbox.line.color.rgb = C_BORDER
        sbox.line.width = Pt(1)

        stb = s2.shapes.add_textbox(dia_x + Inches(0.35), sy + Inches(0.08), src_w - Inches(0.2), Inches(0.55))
        stf = stb.text_frame
        p = stf.paragraphs[0]
        p.text = f"SOURCE 0{i+1}: {src['name'].upper()}"
        p.font.name = FONT_MONO
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_TEAL
        p = stf.add_paragraph()
        p.text = src["detail"]
        p.font.name = FONT_HEAD
        p.font.size = Pt(11)
        p.font.color.rgb = C_INK

    # Center Connector
    tb_arr = s2.shapes.add_textbox(dia_x + Inches(4.7), dia_y + Inches(1.5), Inches(0.8), Inches(1.0))
    p = tb_arr.text_frame.paragraphs[0]
    p.text = "→"
    p.font.name = FONT_HEAD
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # Right: Focal Black Card
    fc_x, fc_y, fc_w, fc_h = dia_x + Inches(5.6), dia_y + Inches(0.25), Inches(6.55), Inches(3.45)
    fc = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fc_x, fc_y, fc_w, fc_h)
    fc.fill.solid()
    fc.fill.fore_color.rgb = C_INK
    fc.line.width = 0

    tb_fc = s2.shapes.add_textbox(fc_x + Inches(0.4), fc_y + Inches(0.35), fc_w - Inches(0.8), fc_h - Inches(0.7))
    tf_fc = tb_fc.text_frame
    tf_fc.word_wrap = True
    p = tf_fc.paragraphs[0]
    p.text = "THE CRITICAL INFORMATION GAP"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xD4, 0x85)
    p.space_after = Pt(12)

    p = tf_fc.add_paragraph()
    p.text = "“What verifiable evidence justifies moving this money?”"
    p.font.name = FONT_HEAD
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(14)

    p = tf_fc.add_paragraph()
    p.text = "Without an auditable evidence pack, the seller either refunds blind (losing COGS and 2-way freight) or disputes unfairly (destroying customer trust and merchant ratings)."
    p.font.name = FONT_HEAD
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(0xB5, 0xCC, 0xC2)

    # Bottom 2 Strips
    b_y = dia_y + dia_h + Inches(0.18)
    b_w = Inches(6.08)

    # Left Bottom Strip
    bs_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), b_y, b_w, Inches(1.0))
    bs_l.fill.solid()
    bs_l.fill.fore_color.rgb = C_PAPER_TINT
    bs_l.line.color.rgb = C_BORDER
    bs_l.line.width = Pt(1)

    tb_bsl = s2.shapes.add_textbox(Inches(0.6), b_y + Inches(0.1), b_w - Inches(0.3), Inches(0.8))
    tf = tb_bsl.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "9% vs 15.14% · Industry Baseline"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf.add_paragraph()
    p.text = "9% NRF 2025 criminal fraud only vs 15.14% Appriss/Deloitte 2024 fraud + policy abuse (definitions differ)."
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_INK_LIGHT

    # Right Bottom Strip
    bs_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.79), b_y, b_w, Inches(1.0))
    bs_r.fill.solid()
    bs_r.fill.fore_color.rgb = C_PAPER_TINT
    bs_r.line.color.rgb = C_BORDER
    bs_r.line.width = Pt(1)

    tb_bsr = s2.shapes.add_textbox(Inches(6.93), b_y + Inches(0.1), b_w - Inches(0.3), Inches(0.8))
    tf = tb_bsr.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Manual Triage · Asymmetric Risk"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_DARK
    p = tf.add_paragraph()
    p.text = "High cognitive overhead on unboxing triage; memory-based manual checks create friction and financial leaks."
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_INK_LIGHT

    add_footer(s2, 2)
    add_notes(s2, d2["speaker_notes"])

    # =========================================================
    # SLIDE 3: STAR DEMO (MASSIVE COMPARISON)
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, C_PAPER)
    d3 = slides_data[3]

    tb = s3.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "STAR DEMO · CASE B · CORE BENCHMARK"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s3.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The order says White. The unboxing video shows Black."
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = "Policy P-03 triggers ESCALATE. No mock label created. No refund draft. No denial without human review."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_MUTED

    # Massive Side-by-side Stage
    stg_w = Inches(5.6)
    stg_h = Inches(4.0)
    stg_y = Inches(1.65)

    # Left Board: Store Manifest
    bd_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), stg_y, stg_w, stg_h)
    bd_l.fill.solid()
    bd_l.fill.fore_color.rgb = C_WHITE
    bd_l.line.color.rgb = C_BORDER_STRONG
    bd_l.line.width = Pt(1)

    tb = s3.shapes.add_textbox(Inches(0.7), stg_y + Inches(0.2), stg_w - Inches(0.5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "ORDER MANIFEST · STORE TRUTH (#ORD-7729)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # Visual container
    vis_l = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), stg_y + Inches(0.65), stg_w - Inches(0.5), Inches(2.1))
    vis_l.fill.solid()
    vis_l.fill.fore_color.rgb = C_PAPER
    vis_l.line.color.rgb = C_BORDER
    vis_l.line.width = Pt(1)

    tb_vis_l = s3.shapes.add_textbox(Inches(0.9), stg_y + Inches(1.1), stg_w - Inches(0.9), Inches(1.2))
    p = tb_vis_l.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "STUDIO WIRELESS HEADPHONES\n[ MATTE WHITE CASING ]\nSKU: HEAD-WHT-02 · ₹1,499"
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_INK

    tb_spec_l = s3.shapes.add_textbox(Inches(0.7), stg_y + Inches(2.9), stg_w - Inches(0.5), Inches(0.9))
    tf = tb_spec_l.text_frame
    p = tf.paragraphs[0]
    p.text = "SKU: HEAD-WHT-02 · Color: Matte White\nPrice: ₹1,499 paid · Delivered: 4d ago (Shippo #SH-9012)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    # Center Mismatch Pill
    pill = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.23), stg_y + Inches(1.4), Inches(0.86), Inches(0.86))
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_CORAL
    pill.line.width = 0
    p = pill.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "≠"
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Right Board: Unboxing Video
    bd_r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.27), stg_y, stg_w, stg_h)
    bd_r.fill.solid()
    bd_r.fill.fore_color.rgb = RGBColor(0xFF, 0xFD, 0xFC)
    bd_r.line.color.rgb = C_CORAL_BORDER
    bd_r.line.width = Pt(1.5)

    tb = s3.shapes.add_textbox(Inches(7.5), stg_y + Inches(0.2), stg_w - Inches(0.5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "MOCK EVIDENCE PACK · VIDEO CITATION (FRAME 00:07)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_CORAL

    vis_r = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.5), stg_y + Inches(0.65), stg_w - Inches(0.5), Inches(2.1))
    vis_r.fill.solid()
    vis_r.fill.fore_color.rgb = RGBColor(0x18, 0x22, 0x1E)
    vis_r.line.width = 0

    tb_vis_r = s3.shapes.add_textbox(Inches(7.7), stg_y + Inches(1.05), stg_w - Inches(0.9), Inches(1.3))
    p = tb_vis_r.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "CUSTOMER UNBOXING FOOTAGE\n[ MIDNIGHT BLACK CASING DETECTED ]\nREC 00:07 / 00:24 · Lighting: 142/255"
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xF7, 0xF5, 0xF0)

    tb_spec_r = s3.shapes.add_textbox(Inches(7.5), stg_y + Inches(2.9), stg_w - Inches(0.5), Inches(0.9))
    tf = tb_spec_r.text_frame
    p = tf.paragraphs[0]
    p.text = "Observed: HEAD-BLK-09 (Midnight Black)\nExact Barcode: Unverified · P-03 Color Match: FAIL"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_CORAL_DARK

    # Bottom Decision Strip
    dec_y = stg_y + stg_h + Inches(0.18)
    dec = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), dec_y, Inches(12.41), Inches(1.0))
    dec.fill.solid()
    dec.fill.fore_color.rgb = C_WHITE
    dec.line.color.rgb = C_BORDER_STRONG
    dec.line.width = Pt(1)

    tb_dec = s3.shapes.add_textbox(Inches(0.65), dec_y + Inches(0.1), Inches(12.0), Inches(0.8))
    tf = tb_dec.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "POLICY P-03 RESULT: ESCALATE TO SELLER"
    p.font.name = FONT_MONO
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_CORAL_DARK
    p = tf.add_paragraph()
    p.text = "3 Preserved Seller Actions: (1) Accept as warehouse mis-shipment · (2) Request photo of outside-box shipping label · (3) Hold return for physical parcel inspection.\nServer Cap Formula: 0 < amount <= min(₹4,999, remaining_refundable, order_paid)"
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_INK_LIGHT

    add_footer(s3, 3)
    add_notes(s3, d3["speaker_notes"])

    # =========================================================
    # SLIDE 4: SYSTEM ARCHITECTURE (DARK)
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, C_DARK_BG)
    d4 = slides_data[4]

    tb = s4.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "SYSTEM ARCHITECTURE · FAIL-CLOSED GOVERNANCE"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT

    tb = s4.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Models observe. Code verifies. Seller authorizes."
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p = tf.add_paragraph()
    p.text = "Exactly 3 base model calls per case. Weak evidence escalates. Never releases automated refunds."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_DARK_MUTED

    # Top Pipeline Bar
    pip_y = Inches(1.65)
    pip = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), pip_y, Inches(12.41), Inches(1.05))
    pip.fill.solid()
    pip.fill.fore_color.rgb = C_DARK_CARD
    pip.line.color.rgb = C_DARK_BORDER
    pip.line.width = Pt(1)

    steps = [
        ("1", "Planner", "Typed actions"),
        ("2", "Tools", "FFmpeg + APIs"),
        ("3", "Policy", "Rules P-01–08"),
        ("4", "Critic", "Review cited frames"),
        ("5", "Seller Gate", "1-Click auth"),
        ("6", "Audit Log", "Append-only JSON")
    ]
    for idx, (num, name, desc) in enumerate(steps):
        step_x = Inches(0.65 + idx * 2.05)
        tb_step = s4.shapes.add_textbox(step_x, pip_y + Inches(0.12), Inches(1.85), Inches(0.8))
        tf = tb_step.text_frame
        p = tf.paragraphs[0]
        p.text = f"{num}. {name}"
        p.font.name = FONT_MONO
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_ACCENT if num == "5" else C_DARK_INK
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT_HEAD
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK_MUTED

    # 3 Guarantee Cards
    g_y = Inches(2.9)
    g_w = Inches(3.95)
    g_h = Inches(2.75)

    # Card 1: Budget
    gc1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), g_y, g_w, g_h)
    gc1.fill.solid()
    gc1.fill.fore_color.rgb = C_DARK_CARD
    gc1.line.color.rgb = C_DARK_BORDER
    gc1.line.width = Pt(1)

    tb = s4.shapes.add_textbox(Inches(0.65), g_y + Inches(0.2), g_w - Inches(0.4), g_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CALL BUDGET"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Exactly 3 Model Calls"
    p.font.name = FONT_HEAD
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "1. Planner proposes typed tool actions.\n2. Vision model inspects server-generated keyframes.\n3. Critic evaluates citations against policy.\n\nModel never accesses payment APIs directly."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK_MUTED

    # Card 2: Re-film loop
    gc2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.69), g_y, g_w, g_h)
    gc2.fill.solid()
    gc2.fill.fore_color.rgb = RGBColor(0x14, 0x28, 0x22)
    gc2.line.color.rgb = C_TEAL
    gc2.line.width = Pt(1.5)

    tb = s4.shapes.add_textbox(Inches(4.88), g_y + Inches(0.2), g_w - Inches(0.4), g_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "STATEFUL RESUMPTION"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Case C: Re-Film Loop"
    p.font.name = FONT_HEAD
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "When video is too dark (brightness 28, 2 frames), agent calls ask_for_evidence() with a re-film checklist.\n\nCase state saved to WAITING_FOR_EVIDENCE and resumes on upload (1 retry limit)."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK_MUTED

    # Card 3: Tiny route
    gc3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.92), g_y, g_w, g_h)
    gc3.fill.solid()
    gc3.fill.fore_color.rgb = C_DARK_CARD
    gc3.line.color.rgb = C_DARK_BORDER
    gc3.line.width = Pt(1)

    tb = s4.shapes.add_textbox(Inches(9.11), g_y + Inches(0.2), g_w - Inches(0.4), g_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SECONDARY ROUTE"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf.add_paragraph()
    p.text = "Tiny PS-04 Order Route"
    p.font.name = FONT_HEAD
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_DARK_INK
    p.space_after = Pt(8)
    p = tf.add_paragraph()
    p.text = "For routine \"Where is my order?\" inquiries, agent bypasses vision entirely:\n\nStore API → Courier API → Instant status reply.\n\nZero vision token waste."
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK_MUTED

    # Bottom Banner
    ban_y = g_y + g_h + Inches(0.18)
    ban = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), ban_y, Inches(12.41), Inches(0.75))
    ban.fill.solid()
    ban.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x1D)
    ban.line.color.rgb = C_DARK_BORDER
    ban.line.width = Pt(1)

    tb_ban = s4.shapes.add_textbox(Inches(0.65), ban_y + Inches(0.12), Inches(12.0), Inches(0.5))
    tf = tb_ban.text_frame
    p = tf.paragraphs[0]
    p.text = "FAIL-CLOSED FINANCIAL GATE: requires_seller_approval: true enforced at schema layer.\nServer Cap: 0 < amount <= min(₹4,999, remaining_refundable, order_paid)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10.5)
    p.font.color.rgb = RGBColor(0xFF, 0xD4, 0x85)

    add_footer(s4, 4, dark=True)
    add_notes(s4, d4["speaker_notes"])

    # =========================================================
    # SLIDE 5: 4 BENCHMARK CASES (ASYMMETRIC SPLIT)
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, C_PAPER)
    d5 = slides_data[5]

    tb = s5.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "EVALUATION SUITE · 4 CONTROLLED SCENARIOS"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s5.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Different evidence. Controlled, predictable action."
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = "Every case produces an auditable mock evidence pack. Zero hallucinated policy approvals."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_MUTED

    # Left: Case B Hero Card
    b_x, b_y, b_w, b_h = Inches(0.46), Inches(1.65), Inches(5.4), Inches(5.1)
    bc = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_x, b_y, b_w, b_h)
    bc.fill.solid()
    bc.fill.fore_color.rgb = C_WHITE
    bc.line.color.rgb = C_CORAL_BORDER
    bc.line.width = Pt(2)

    tb_bc = s5.shapes.add_textbox(b_x + Inches(0.3), b_y + Inches(0.25), b_w - Inches(0.6), b_h - Inches(0.5))
    tf = tb_bc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CORE BENCHMARK SCENARIO · ESCALATE"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_CORAL
    p = tf.add_paragraph()
    p.text = "Case B: Color Contradiction"
    p.font.name = FONT_HEAD
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p.space_after = Pt(12)

    p = tf.add_paragraph()
    p.text = "Order Record: HEAD-WHT-02 (White, ₹1,499)\nObserved Frame 00:07: HEAD-BLK-09 (Midnight Black)\nPolicy Evaluation: P-03 Color Match: FAIL\nAutomated Action: No label · No refund draft · No denial"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_CORAL_DARK
    p.space_after = Pt(16)

    p = tf.add_paragraph()
    p.text = "PRESERVED SELLER OPTIONS IN MOCK EVIDENCE PACK:\n• Accept as warehouse mis-shipment\n• Request customer photo of outer box shipping barcode\n• Hold case for physical receipt inspection at depot"
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.color.rgb = C_INK_LIGHT

    # Right: Compact Stack of Cases A, C, D
    rw_x = Inches(6.08)
    rw_w = Inches(6.79)
    rw_h = Inches(1.58)

    cases_right = [
        ("Case A: Valid Return (ORD-1042)", "PREPARE SELLER REVIEW", C_TEAL_DARK, C_TEAL_LIGHT,
         "Evidence: ORD-1042 · ₹1,499 · SKU matches catalog · Tag clearly present · Within 7-day return window.",
         "Outcome: Mock shipping label created · Refund draft placed in pending_return_receipt (not executed)."),
        ("Case C: Low-Light Video (Re-Film Loop)", "ASK & RESUME", C_AMBER_DARK, C_AMBER_LIGHT,
         "Evidence: 4-second dark unboxing video · Brightness score 28/255 · Only 2 usable keyframes extracted.",
         "Outcome: Agent issues re-film checklist · Case stays in WAITING_FOR_EVIDENCE · Resumes on upload (1 retry limit)."),
        ("Case D: Continuous Unboxing Sequence", "INSPECT SEQUENCE", C_INK, C_PAPER_TINT,
         "Evidence: Single continuous take from opening sealed parcel to displaying label in one sequence.",
         "Outcome: Inspects sequence continuity and flags abrupt transitions. Does not claim proof of identical item.")
    ]

    for idx, (title, badge, b_col, bg_col, ev, out) in enumerate(cases_right):
        cy = b_y + Inches(idx * 1.76)
        cr = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rw_x, cy, rw_w, rw_h)
        cr.fill.solid()
        cr.fill.fore_color.rgb = C_WHITE
        cr.line.color.rgb = C_BORDER
        cr.line.width = Pt(1)

        tb = s5.shapes.add_textbox(rw_x + Inches(0.25), cy + Inches(0.12), rw_w - Inches(0.5), rw_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}   [{badge}]"
        p.font.name = FONT_HEAD
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = b_col
        p = tf.add_paragraph()
        p.text = ev
        p.font.name = FONT_HEAD
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_MUTED
        p = tf.add_paragraph()
        p.text = out
        p.font.name = FONT_MONO
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_TEAL_DARK

    add_footer(s5, 5)
    add_notes(s5, d5["speaker_notes"])

    # =========================================================
    # SLIDE 6: 24-HOUR PROTOTYPE PLAN
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, C_PAPER)
    d6 = slides_data[6]

    tb = s6.shapes.add_textbox(Inches(0.46), Inches(0.36), Inches(12.0), Inches(0.25))
    p = tb.text_frame.paragraphs[0]
    p.text = "24-HOUR PROTOTYPE PLAN · TEAM ROOKIE (3 MEMBERS)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    tb = s6.shapes.add_textbox(Inches(0.46), Inches(0.65), Inches(12.4), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "A credible, buildable prototype in 24 hours."
    p.font.name = FONT_HEAD
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = tf.add_paragraph()
    p.text = "Strict scope discipline, mock data fixtures, and single-threaded ownership ensure on-time delivery."
    p.font.name = FONT_HEAD
    p.font.size = Pt(14)
    p.font.color.rgb = C_MUTED

    # Full Width Timeline Bar
    t_y = Inches(1.65)
    t_bar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), t_y, Inches(12.41), Inches(1.2))
    t_bar.fill.solid()
    t_bar.fill.fore_color.rgb = C_WHITE
    t_bar.line.color.rgb = C_BORDER_STRONG
    t_bar.line.width = Pt(1)

    t_slots = [
        ("00 – 06h", "Foundation", "Mock Store/Shippo APIs · LangGraph state schema · Policy P-01 to P-08"),
        ("06 – 12h", "Evidence Pipeline", "FFmpeg ~1fps extractor · Vision prompt schemas · Seller Console UI mock"),
        ("12 – 18h", "Integration", "Case C pause/resume state loop · 4 demo cases (A–D) · Critic review"),
        ("18 – 24h", "Validation", "Fail-closed authorization tests · Boundary QA · Video demo recording")
    ]
    for idx, (hrs, phase, desc) in enumerate(t_slots):
        tx = Inches(0.65 + idx * 3.08)
        tb = s6.shapes.add_textbox(tx, t_y + Inches(0.12), Inches(2.9), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{hrs} · {phase}"
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TEAL
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = FONT_HEAD
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED

    # People 3-Column Strip
    p_y = Inches(3.05)
    p_w = Inches(3.95)
    p_h = Inches(2.35)

    people = [
        ("PERSON A", "Backend, Policy & Audit",
         "LangGraph state machine, policy engine P-01 to P-08, Critic validation, append-only JSON audit trail, server-side refund cap calculation."),
        ("PERSON B", "Video & Seller Console",
         "FFmpeg keyframe pipeline (~1fps sampling), vision model structured citation prompt, interactive 1-click Seller Console mockup."),
        ("PERSON C", "Integration, Scenarios & QA",
         "Test fixtures for Cases A–D, Case C stateful pause/resume lifecycle, fail-closed human gate testing, benchmark runner, video capture.")
    ]

    for idx, (role, name, tasks) in enumerate(people):
        px = Inches(0.46 + idx * 4.23)
        pc = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, p_y, p_w, p_h)
        pc.fill.solid()
        pc.fill.fore_color.rgb = C_PAPER
        pc.line.color.rgb = C_BORDER
        pc.line.width = Pt(1)

        tb = s6.shapes.add_textbox(px + Inches(0.2), p_y + Inches(0.18), p_w - Inches(0.4), p_h - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = role
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TEAL
        p = tf.add_paragraph()
        p.text = name
        p.font.name = FONT_HEAD
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_INK
        p.space_after = Pt(8)
        p = tf.add_paragraph()
        p.text = tasks
        p.font.name = FONT_HEAD
        p.font.size = Pt(11)
        p.font.color.rgb = C_MUTED

    # Bottom Scope Discipline Band
    sc_y = Inches(5.6)
    sc = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.46), sc_y, Inches(12.41), Inches(1.15))
    sc.fill.solid()
    sc.fill.fore_color.rgb = C_INK
    sc.line.width = 0

    tb_sc1 = s6.shapes.add_textbox(Inches(0.65), sc_y + Inches(0.15), Inches(4.5), Inches(0.85))
    tf1 = tb_sc1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "MUST DELIVER (CORE PROTOTYPE)"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xD4, 0x85)
    p = tf1.add_paragraph()
    p.text = "Mock APIs · Policy engine · Case B mismatch · Case C re-film · Seller gate · Tiny order route"
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = RGBColor(0xDF, 0xEA, 0xE5)

    tb_sc2 = s6.shapes.add_textbox(Inches(5.4), sc_y + Inches(0.15), Inches(4.2), Inches(0.85))
    tf2 = tb_sc2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "CUT FIRST IF BEHIND"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0x8C, 0x69)
    p = tf2.add_paragraph()
    p.text = "WhatsApp bot · Exchanges · Multi-seller auth · Audio analysis · Public cloud deployment"
    p.font.name = FONT_HEAD
    p.font.size = Pt(10.5)
    p.font.color.rgb = RGBColor(0xDF, 0xEA, 0xE5)

    tb_sc3 = s6.shapes.add_textbox(Inches(9.8), sc_y + Inches(0.15), Inches(2.9), Inches(0.85))
    tf3 = tb_sc3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "FINANCIAL CEILING FORMULA"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_TEAL_ACCENT
    p = tf3.add_paragraph()
    p.text = "0 < amt <= min(₹4,999, rem_ref, paid)"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE

    add_footer(s6, 6)
    add_notes(s6, d6["speaker_notes"])

    # Save PPTX Deliverable
    pptx_deliv = deliverables_dir / "ReturnLens_v5.pptx"
    pptx_root = base_dir / "ReturnLens_v5.pptx"
    prs.save(str(pptx_deliv))
    prs.save(str(pptx_root))
    print(f"Generated ReturnLens_v5.pptx successfully: {pptx_deliv} ({os.path.getsize(pptx_deliv)} bytes)")

if __name__ == "__main__":
    build_v5_pptx()
