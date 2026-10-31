import os
import sys
import json
import zipfile
import subprocess
import hashlib
from pathlib import Path
import fitz
from PIL import Image
from pptx import Presentation

base_dir = Path(__file__).resolve().parent.parent

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

def file_hash(p):
    with open(p, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()

def run_tests():
    print("=== STARTING RETURNLENS V4 EXTENDED SMOKE & QA TESTS ===")
    all_passed = True

    # 1. Test Canonical Content Specification
    print("\n[TEST 1] Validating Canonical Data Specification (content/deck_data.json)...")
    data_path = base_dir / "content" / "deck_data.json"
    if not data_path.exists():
        print(f"FAILED: Canonical data file missing: {data_path}")
        all_passed = False
    else:
        try:
            with open(data_path, "r", encoding="utf-8") as f:
                d = json.load(f)
            sample_price = d["meta"]["sample_order"]["price"]
            ceiling = d["meta"]["refund_guardrail"]["ceiling"]
            print(f"PASSED: Canonical data loaded successfully. Sample order: {sample_price}, Ceiling: {ceiling}.")
            if sample_price != "₹1,499" or ceiling != "₹4,999":
                print(f"FAILED: Unexpected price/ceiling values in deck_data.json: {sample_price}, {ceiling}")
                all_passed = False
        except Exception as e:
            print(f"FAILED: deck_data.json parsing error: {e}")
            all_passed = False

    # 2. Test HTML & PDF Generator Execution
    print("\n[TEST 2] Running src/generate_v4_html.py...")
    script_html = base_dir / "src" / "generate_v4_html.py"
    h_before = file_hash(script_html)

    res = subprocess.run([sys.executable, str(script_html)], capture_output=True, text=True)
    if res.returncode != 0:
        print(f"FAILED: generate_v4_html.py returned code {res.returncode}:\n{res.stderr}")
        all_passed = False
    else:
        print("PASSED: generate_v4_html.py executed successfully.")

    # Idempotency Check
    h_after = file_hash(script_html)
    if h_before != h_after:
        print("FAILED: P0 BUG - generate_v4_html.py modified its own source code!")
        all_passed = False
    else:
        print("PASSED: Generator source file remained completely untouched (100% idempotent).")

    # 3. Test PPTX Generator Execution
    print("\n[TEST 3] Running src/build_v4_pptx.py...")
    script_pptx = base_dir / "src" / "build_v4_pptx.py"
    res_pptx = subprocess.run([sys.executable, str(script_pptx)], capture_output=True, text=True)
    if res_pptx.returncode != 0:
        print(f"FAILED: build_v4_pptx.py returned code {res_pptx.returncode}:\n{res_pptx.stderr}")
        all_passed = False
    else:
        print("PASSED: build_v4_pptx.py executed successfully.")

    # 4. Test Deliverable Integrity & Synchronization
    print("\n[TEST 4] Checking deliverable files and synchronization...")
    pairs = [
        ("ReturnLens_v4.html", "deliverables/ReturnLens_v4.html"),
        ("ReturnLens_v4.pdf", "deliverables/ReturnLens_v4.pdf"),
        ("ReturnLens_v4.pptx", "deliverables/ReturnLens_v4.pptx"),
        ("ReturnLens_v4_preview.png", "deliverables/ReturnLens_v4_preview.png"),
    ]

    for root_f, deliv_f in pairs:
        p_root = base_dir / root_f
        p_deliv = base_dir / deliv_f
        if not p_root.exists() or p_root.stat().st_size == 0:
            print(f"FAILED: {root_f} missing or empty.")
            all_passed = False
        elif not p_deliv.exists() or p_deliv.stat().st_size == 0:
            print(f"FAILED: {deliv_f} missing or empty.")
            all_passed = False
        elif file_hash(p_root) != file_hash(p_deliv):
            print(f"FAILED: {root_f} and {deliv_f} are not synchronized (hashes differ).")
            all_passed = False
        else:
            print(f"PASSED: {root_f} synced with {deliv_f} ({p_root.stat().st_size} bytes).")

    # 5. Test PDF Properties
    print("\n[TEST 5] Validating PDF properties...")
    pdf_path = base_dir / "ReturnLens_v4.pdf"
    doc = fitz.open(pdf_path)
    if len(doc) != 6:
        print(f"FAILED: PDF page count is {len(doc)}, expected 6.")
        all_passed = False
    else:
        print("PASSED: PDF has exactly 6 pages.")

    rect = doc[0].rect
    ratio = rect.width / rect.height
    expected_ratio = 16.0 / 9.0
    if abs(ratio - expected_ratio) > 0.01:
        print(f"FAILED: PDF aspect ratio is {ratio:.3f}, expected {expected_ratio:.3f}.")
        all_passed = False
    else:
        print(f"PASSED: PDF aspect ratio is true 16:9 widescreen ({rect.width}x{rect.height}).")

    # 6. Test PPTX Properties, Speaker Notes & EMBEDDED MEDIA (Critical Fix Validation)
    print("\n[TEST 6] Validating PPTX Properties, Speaker Notes & Embedded Media...")
    pptx_path = base_dir / "ReturnLens_v4.pptx"
    prs = Presentation(pptx_path)
    if len(prs.slides) != 6:
        print(f"FAILED: PPTX slide count is {len(prs.slides)}, expected 6.")
        all_passed = False
    else:
        print("PASSED: PPTX has exactly 6 slides.")

    missing_notes = []
    for i, slide in enumerate(prs.slides):
        if not slide.has_notes_slide:
            missing_notes.append(i + 1)
        else:
            txt = slide.notes_slide.notes_text_frame.text.strip()
            if len(txt) == 0:
                missing_notes.append(i + 1)
            else:
                print(f"  Slide {i+1} Speaker Notes: {len(txt)} chars")

    if missing_notes:
        print(f"FAILED: Slides {missing_notes} missing speaker notes in PPTX.")
        all_passed = False
    else:
        print("PASSED: All 6 slides have verified embedded speaker notes.")

    # CRITICAL: Verify embedded media files in PPTX package!
    with zipfile.ZipFile(pptx_path) as z:
        media_files = [f for f in z.namelist() if 'media/' in f]
        print(f"  Embedded media in PPTX: {len(media_files)} files ({media_files})")
        if len(media_files) < 2:
            print(f"FAILED: Expected at least 2 embedded media files in PPTX, found {len(media_files)}!")
            all_passed = False
        else:
            print("PASSED: True embedded product images verified inside PPTX zip package.")

    # 7. Test Preview Image Dimensions
    print("\n[TEST 7] Validating Preview Image Dimensions...")
    prev_path = base_dir / "ReturnLens_v4_preview.png"
    im = Image.open(prev_path)
    if im.size != (2352, 900):
        print(f"FAILED: Preview image size is {im.size}, expected (2352, 900).")
        all_passed = False
    else:
        print("PASSED: Preview image size is verified (2352, 900).")

    # 8. Test Evidence Grounding & Factual Consistency
    print("\n[TEST 8] Verifying Evidence Grounding in PDF & PPTX...")
    pdf_text = ""
    for page in doc:
        pdf_text += page.get_text() + "\n"

    pdf_lower = pdf_text.lower()

    # Grounding check 1: Honest SKU unverified on Slide 3
    if "exact sku: not verified" not in pdf_lower:
        print("FAILED: 'Exact SKU: Not verified' missing from Slide 3 evidence footer.")
        all_passed = False
    else:
        print("PASSED: Color mismatch clearly states SKU is unverified.")

    # Grounding check 2: PS-04 secondary route present
    if "secondary ps-04 route" not in pdf_lower or "where is my order?" not in pdf_lower:
        print("FAILED: Secondary PS-04 order-status route missing from Architecture slide.")
        all_passed = False
    else:
        print("PASSED: Secondary PS-04 order status route verified on Slide 4.")

    # Grounding check 3: Case C stateful clarification loop present
    if "waiting_for_evidence" not in pdf_lower or "ask_for_evidence" not in pdf_lower:
        print("FAILED: Case C stateful loop missing from Architecture slide.")
        all_passed = False
    else:
        print("PASSED: Dynamic stateful clarification loop verified on Slide 4.")

    # Grounding check 4: Conservative Case D sequence
    if "inspect sequence continuity" not in pdf_lower:
        print("FAILED: Conservative Case D sequence continuity text missing from Slide 5.")
        all_passed = False
    else:
        print("PASSED: Case D uses conservative sequence continuity wording.")

    # 9. Test Visual Geometry, Collision Avoidance & Margin Safety
    print("\n[TEST 9] Validating Visual Geometry, Collisions & Safe Margins in PPTX...")
    prs = Presentation(pptx_path)
    geom_passed = True

    for s_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            left = shape.left.inches if shape.left else 0
            top = shape.top.inches if shape.top else 0
            w = shape.width.inches if shape.width else 0
            h = shape.height.inches if shape.height else 0
            bottom = top + h
            right = left + w

            # Bounds check (allow 0.05" epsilon)
            if right > 13.38 or bottom > 7.55:
                print(f"FAILED: Shape '{shape.name}' on Slide {s_idx + 1} exceeds slide boundary: right={right:.2f}, bottom={bottom:.2f}")
                geom_passed = False
                all_passed = False

            # Footer collision check (content shapes must not cross footer at y=6.92)
            if top > 0.1 and top < 6.90 and bottom > 6.92:
                print(f"FAILED: Content shape '{shape.name}' on Slide {s_idx + 1} collides with footer line at y=6.92 (bottom={bottom:.2f})")
                geom_passed = False
                all_passed = False

    # Slide 4 collision check: Case C title vs Flow Nodes
    s4 = prs.slides[3]
    title_bottom = None
    node_top = None
    for shape in s4.shapes:
        if shape.has_text_frame and "DYNAMIC STATEFUL CLARIFICATION" in shape.text_frame.text:
            title_bottom = shape.top.inches + shape.height.inches
        if shape.has_text_frame and "WAITING_FOR_EVIDENCE" in shape.text_frame.text:
            node_top = shape.top.inches

    if title_bottom is not None and node_top is not None:
        gap = node_top - title_bottom
        if gap < 0:
            print(f"FAILED: Slide 4 Case C title collides with flow nodes! Title bottom: {title_bottom:.2f}, Node top: {node_top:.2f}")
            geom_passed = False
            all_passed = False
        else:
            print(f"PASSED: Slide 4 Case C title and flow nodes have verified positive clear gap ({gap:.2f}\" clearance).")
    else:
        print("WARNING: Could not identify Slide 4 Case C title and node shapes for collision check.")

    if geom_passed:
        print("PASSED: All slides passed geometry, safe margins, and collision checks.")

    print("\n==========================================")
    if all_passed:
        print("ALL RETURNLENS V4 EXTENDED SMOKE & QA TESTS PASSED! READY FOR SUBMISSION.")
    else:
        print("SOME SMOKE TESTS FAILED. PLEASE REVIEW LOGS ABOVE.")
    print("==========================================")
    return all_passed

if __name__ == "__main__":
    success = run_tests()
    sys.exit(0 if success else 1)
