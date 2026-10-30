import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_v3_deck(output_pptx):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Redesigned Color Palette
    COLOR_BG_LIGHT = RGBColor(245, 243, 239)   # #F5F3EF
    COLOR_WHITE = RGBColor(255, 255, 255)      # #FFFFFF
    COLOR_TEXT_DARK = RGBColor(23, 38, 33)     # #172621
    COLOR_TEXT_MUTED = RGBColor(90, 102, 97)   # #5A6661
    COLOR_LINE = RGBColor(226, 223, 216)       # #E2DFD8
    COLOR_TEAL = RGBColor(0, 107, 85)          # #006B55
    COLOR_AMBER = RGBColor(216, 140, 49)       # #D88C31
    COLOR_RED = RGBColor(197, 57, 41)          # #C53929

    # Dark Theme Palette (Slide 4)
    COLOR_DARK_BG = RGBColor(13, 24, 20)       # #0D1814
    COLOR_DARK_CARD = RGBColor(20, 34, 29)     # #14221D
    COLOR_DARK_LINE = RGBColor(31, 52, 44)     # #1F342C
    COLOR_TEXT_LIGHT = RGBColor(237, 235, 230) # #EDEBE6
    COLOR_TEXT_LIGHT_MUTED = RGBColor(149, 163, 157) # #95A39D
    COLOR_TEAL_LIGHT = RGBColor(62, 201, 182)  # #3EC9B6

    img_white_path = os.path.abspath(r"D:\Code\Code\AIO\Code\ZaloAIC\assets_v3\product_white_headphones.png")
    img_black_path = os.path.abspath(r"D:\Code\Code\AIO\Code\ZaloAIC\assets_v3\video_frame_black_headphones.png")

    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, sub_text=None, is_dark=False):
        tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.733), Inches(0.28))
        tf_tag = tb_tag.text_frame
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_TEAL_LIGHT if is_dark else COLOR_TEAL

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.6))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE if is_dark else COLOR_TEXT_DARK

        if sub_text:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.32), Inches(11.733), Inches(0.35))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = sub_text
            p_sub.font.name = "Arial"
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = COLOR_TEXT_LIGHT_MUTED if is_dark else COLOR_TEXT_MUTED

    def add_footer(slide, text_left, page_num, is_dark=False):
        y_pos = Inches(7.0)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_DARK_LINE if is_dark else COLOR_LINE
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), y_pos + Inches(0.08), Inches(11.733), Inches(0.3))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"{text_left}  ·  Slide {page_num} / 6"
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_LIGHT_MUTED if is_dark else COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 1: Product Introduction
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, COLOR_BG_LIGHT)
    add_footer(s1, "ReturnLens · Evidence-Backed Return Triage", 1)

    # Top tag
    tb_tag = s1.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.3))
    tf_tag = tb_tag.text_frame
    p = tf_tag.paragraphs[0]
    p.text = "AGENTX HACKATHON 2026 · PS-04 TRACK"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    # Left text
    tb_l = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6.0), Inches(5.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

    p = tf_l.paragraphs[0]
    p.text = "ReturnLens"
    p.font.name = "Arial"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_l.add_paragraph()
    p.space_before = Pt(12)
    p.text = "Review returns with cited video evidence."
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    p = tf_l.add_paragraph()
    p.space_before = Pt(16)
    p.text = "A focused return triage agent for solo direct-to-consumer sellers. It links buyer unboxing clips to store records and policy rules for 1-click decision making."
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_l.add_paragraph()
    p.space_before = Pt(24)
    p.text = "Team Rookie  ·  3 Members  ·  24-Hour Prototype  ·  Mock Data & Staged Evidence"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_DARK

    # Right Product Card
    rx = Inches(7.2)
    ry = Inches(1.1)
    rw = Inches(5.3)
    rh = Inches(5.5)

    card1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, ry, rw, rh)
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_WHITE
    card1.line.color.rgb = COLOR_LINE
    card1.line.width = Pt(1)

    tb_c1 = s1.shapes.add_textbox(rx + Inches(0.3), ry + Inches(0.25), rw - Inches(0.6), rh - Inches(0.5))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = tf_c1.margin_right = tf_c1.margin_bottom = 0

    p = tf_c1.paragraphs[0]
    p.text = "CASE #RET-2026-084  [PENDING REVIEW]"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    p = tf_c1.add_paragraph()
    p.space_before = Pt(4)
    p.text = "Order #ORD-1042  ·  Customer: Anita S."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_c1.add_paragraph()
    p.space_before = Pt(10)
    p.text = "ITEM: Wireless Over-Ear Headphones (White)\nSKU: HEAD-WHT-02  ·  Paid: ₹1,499  ·  Delivered 4d ago"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_DARK

    # Add Thumbnail
    if os.path.exists(img_white_path):
        s1.shapes.add_picture(img_white_path, rx + Inches(0.3), ry + Inches(1.6), width=Inches(1.8), height=Inches(1.2))

    tb_c1_sub = s1.shapes.add_textbox(rx + Inches(2.2), ry + Inches(1.6), Inches(2.8), Inches(1.2))
    tf_sub = tb_c1_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p = tf_sub.paragraphs[0]
    p.text = "STORE RECORD:\nDelivered in 7d window\nPackaging verified\nPolicy eligible: Yes"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEAL

    # Video proof
    tb_c1_bot = s1.shapes.add_textbox(rx + Inches(0.3), ry + Inches(3.0), rw - Inches(0.6), Inches(2.2))
    tf_bot = tb_c1_bot.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_top = tf_bot.margin_right = tf_bot.margin_bottom = 0

    p = tf_bot.paragraphs[0]
    p.text = "CITED FRAME: 00:14 / 00:45"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_bot.add_paragraph()
    p.space_before = Pt(3)
    p.text = "\"Seam tear on left casing verified. Brand tag matches store SKU.\""
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_bot.add_paragraph()
    p.space_before = Pt(14)
    p.text = "NEXT ACTION:  [ Approve Return (1-Click) ]"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    p = tf_bot.add_paragraph()
    p.space_before = Pt(4)
    p.text = "Seller approval mandatory · Creates mock label & refund draft in pending_return_receipt"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 2: The Seller's Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, COLOR_BG_LIGHT)
    add_header(s2, "The Problem  ·  Solo Seller Bottleneck",
               "Return evidence is scattered across four sources",
               "Sellers make high-stakes return decisions from memory while toggling between fragmented tabs.")
    add_footer(s2, "ReturnLens · Problem Context", 2)

    tabs = [
        ("1. Buyer Chat / Video", "\"Received damaged unit. Need full refund.\"\n📎 unboxing_clip_32s.mp4"),
        ("2. Order Record", "ORD-8492  ·  ₹1,499\nSKU: HEAD-WHT-02\nCustomer: Rohit K."),
        ("3. Delivery Status", "Delivered: 26 Oct (4d ago)\nTracking: TRK-98214-DEL\nCarrier: MockPost"),
        ("4. Return Policy", "7-day limit from delivery.\nTags intact. Original box.\nRefund cap: ₹4,999.")
    ]
    for i, (thead, tbody) in enumerate(tabs):
        tx = Inches(0.8 + i * 2.95)
        ty = Inches(1.85)
        tw = Inches(2.8)
        th = Inches(2.2)

        card = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, ty, tw, th)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_LINE
        card.line.width = Pt(1)

        tb = s2.shapes.add_textbox(tx + Inches(0.2), ty + Inches(0.2), tw - Inches(0.4), th - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = thead.upper()
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL

        p = tf.add_paragraph()
        p.space_before = Pt(8)
        p.text = tbody
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Row
    bx1 = Inches(0.8)
    by1 = Inches(4.35)
    bw1 = Inches(7.0)
    bh1 = Inches(2.3)

    card_cost = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx1, by1, bw1, bh1)
    card_cost.fill.solid()
    card_cost.fill.fore_color.rgb = COLOR_WHITE
    card_cost.line.color.rgb = COLOR_AMBER
    card_cost.line.width = Pt(1.5)

    tb_cost = s2.shapes.add_textbox(bx1 + Inches(0.25), by1 + Inches(0.2), bw1 - Inches(0.5), bh1 - Inches(0.4))
    tf_cost = tb_cost.text_frame
    tf_cost.word_wrap = True
    tf_cost.margin_left = tf_cost.margin_top = tf_cost.margin_right = tf_cost.margin_bottom = 0

    p = tf_cost.paragraphs[0]
    p.text = "THE COST OF GUESSWORK"
    p.font.name = "Arial"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    p = tf_cost.add_paragraph()
    p.space_before = Pt(6)
    p.text = "• Wrong Approval: Unrecoverable cash, shipping fees, and inventory lost to fraud or abuse.\n• Unfair Denial: Permanent loss of customer goodwill, chargebacks, and marketplace penalties."
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_DARK

    bx2 = Inches(8.1)
    by2 = Inches(4.35)
    bw2 = Inches(4.4)
    bh2 = Inches(2.3)

    card_stat = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx2, by2, bw2, bh2)
    card_stat.fill.solid()
    card_stat.fill.fore_color.rgb = COLOR_WHITE
    card_stat.line.color.rgb = COLOR_LINE
    card_stat.line.width = Pt(1)

    tb_stat = s2.shapes.add_textbox(bx2 + Inches(0.25), by2 + Inches(0.2), bw2 - Inches(0.5), bh2 - Inches(0.4))
    tf_stat = tb_stat.text_frame
    tf_stat.word_wrap = True
    tf_stat.margin_left = tf_stat.margin_top = tf_stat.margin_right = tf_stat.margin_bottom = 0

    p = tf_stat.paragraphs[0]
    p.text = "SOURCED INDUSTRY BENCHMARK"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    p = tf_stat.add_paragraph()
    p.space_before = Pt(6)
    p.text = "Fraud estimates differ by definition: 9% (NRF 2025, fraud only) vs 15.14% (Appriss/Deloitte 2024, fraud & abuse)."
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_stat.add_paragraph()
    p.space_before = Pt(4)
    p.text = "*U.S. retail estimates; not measured on our target sellers."
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 3: Product Demonstration (STAR SLIDE)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, COLOR_BG_LIGHT)
    add_header(s3, "Product Demonstration  ·  Visual Evidence Benchmark",
               "A color mismatch needs seller review",
               "When video evidence contradicts store truth, ReturnLens surfaces the discrepancy without auto-denying.")
    add_footer(s3, "Illustrative Prototype · Staged Mock Evidence · No Real Customer PII", 3)

    c_w = Inches(3.7)
    c_h = Inches(4.8)
    cy = Inches(1.85)

    # Card 1: Store Truth
    card_st = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), cy, c_w, c_h)
    card_st.fill.solid()
    card_st.fill.fore_color.rgb = COLOR_WHITE
    card_st.line.color.rgb = COLOR_LINE
    card_st.line.width = Pt(1)

    tb_st_head = s3.shapes.add_textbox(Inches(1.0), cy + Inches(0.15), c_w - Inches(0.4), Inches(0.3))
    tf = tb_st_head.text_frame
    p = tf.paragraphs[0]
    p.text = "STORE API RECORD [ORDER TRUTH]"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    if os.path.exists(img_white_path):
        s3.shapes.add_picture(img_white_path, Inches(1.0), cy + Inches(0.5), width=c_w - Inches(0.4), height=Inches(2.2))

    tb_st_txt = s3.shapes.add_textbox(Inches(1.0), cy + Inches(2.8), c_w - Inches(0.4), Inches(1.8))
    tf_txt = tb_st_txt.text_frame
    tf_txt.word_wrap = True
    p = tf_txt.paragraphs[0]
    p.text = "SKU: HEAD-WHT-02\nProduct: Wireless Headset\nOrdered Color: White\nPrice: ₹1,499 · Delivered 4d ago"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_txt.add_paragraph()
    p.space_before = Pt(6)
    p.text = "Customer: Rohit K. · In 7-day window"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # Card 2: Video Evidence
    card_vd = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.8), cy, c_w, c_h)
    card_vd.fill.solid()
    card_vd.fill.fore_color.rgb = COLOR_WHITE
    card_vd.line.color.rgb = COLOR_AMBER
    card_vd.line.width = Pt(2.5)

    tb_vd_head = s3.shapes.add_textbox(Inches(5.0), cy + Inches(0.15), c_w - Inches(0.4), Inches(0.3))
    tf = tb_vd_head.text_frame
    p = tf.paragraphs[0]
    p.text = "VIDEO EVIDENCE [STAGED FRAME 00:07]"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    if os.path.exists(img_black_path):
        s3.shapes.add_picture(img_black_path, Inches(5.0), cy + Inches(0.5), width=c_w - Inches(0.4), height=Inches(2.2))

    tb_vd_txt = s3.shapes.add_textbox(Inches(5.0), cy + Inches(2.8), c_w - Inches(0.4), Inches(1.8))
    tf_vd = tb_vd_txt.text_frame
    tf_vd.word_wrap = True
    p = tf_vd.paragraphs[0]
    p.text = "Observed color: Black\nSKU: Not verified (unreadable)\nTag: Tag not visible in this frame\nDiscrepancy: Color divergence detected"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_DARK

    # Card 3: Verdict & Options
    card_vr = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), cy, c_w, c_h)
    card_vr.fill.solid()
    card_vr.fill.fore_color.rgb = COLOR_WHITE
    card_vr.line.color.rgb = COLOR_LINE
    card_vr.line.width = Pt(1)

    tb_vr = s3.shapes.add_textbox(Inches(9.0), cy + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
    tf_vr = tb_vr.text_frame
    tf_vr.word_wrap = True
    tf_vr.margin_left = tf_vr.margin_top = tf_vr.margin_right = tf_vr.margin_bottom = 0

    p = tf_vr.paragraphs[0]
    p.text = "POLICY P-03: FAIL (MISMATCH)"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED

    p = tf_vr.add_paragraph()
    p.space_before = Pt(8)
    p.text = "STATUS: NEEDS SELLER REVIEW"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    p = tf_vr.add_paragraph()
    p.space_before = Pt(4)
    p.text = "No automatic denial. No return label or refund draft issued at this stage."
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_vr.add_paragraph()
    p.space_before = Pt(12)
    p.text = "Seller Resolution Options:\n1. Review possible warehouse mis-shipment\n2. Request outside-box shipping label photo\n3. Hold for physical inspection upon return"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_vr.add_paragraph()
    p.space_before = Pt(16)
    p.text = "Illustrative prototype · Staged mock evidence"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SLIDE 4: Workflow & Authority (Dark Theme)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, COLOR_DARK_BG)
    add_header(s4, "Workflow & Decision Authority  ·  Deterministic Architecture",
               "Evidence checks and seller approval",
               "A deterministic pipeline where models observe, code verifies, and the seller authorizes.", is_dark=True)
    add_footer(s4, "ReturnLens · System Architecture & Decision Flow", 4, is_dark=True)

    flow_nodes = [
        ("01 · INGEST", "Receive & Retrieve", "Buyer video via mock chat; queries order & courier APIs."),
        ("02 · OBSERVE", "Model Observation", "Models inspect frames; output structured observations."),
        ("03 · CHECK", "Policy Engine (Code)", "Pure code checks 7d window, delivery date, SKU match, refund cap."),
        ("04 · VERIFY", "Independent Critic", "Audits cited frames to confirm observations are supported."),
        ("05 · AUTHORIZE", "Seller Gate", "Seller 1-click approval permits mock label & refund draft.")
    ]
    for i, (ntag, ntitle, ndesc) in enumerate(flow_nodes):
        nx = Inches(0.8 + i * 2.38)
        ny = Inches(1.9)
        nw = Inches(2.25)
        nh = Inches(2.3)

        card = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, nx, ny, nw, nh)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_DARK_CARD
        card.line.color.rgb = COLOR_TEAL_LIGHT if i == 2 else COLOR_DARK_LINE
        card.line.width = Pt(2) if i == 2 else Pt(1)

        tb = s4.shapes.add_textbox(nx + Inches(0.15), ny + Inches(0.15), nw - Inches(0.3), nh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = ntag
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL_LIGHT

        p = tf.add_paragraph()
        p.space_before = Pt(4)
        p.text = ntitle
        p.font.name = "Arial"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

        p = tf.add_paragraph()
        p.space_before = Pt(6)
        p.text = ndesc
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_LIGHT_MUTED

    # 3 Branches
    branches = [
        ("UNCLEAR EVIDENCE", "Asks for specific photo (e.g. tag/serial) and resumes the exact same case.", COLOR_AMBER),
        ("MISMATCH / CONTRADICTION", "Escalates directly to seller review with proof pack; never auto-denies.", COLOR_RED),
        ("ALL POLICY CHECKS PASS", "Prepares 1-click approval package for seller; draft in pending_return_receipt.", COLOR_TEAL_LIGHT)
    ]
    for i, (btag, bdesc, bcol) in enumerate(branches):
        bx = Inches(0.8 + i * 4.0)
        by = Inches(4.4)
        bw = Inches(3.7)
        bh = Inches(1.4)

        card = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, bw, bh)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_DARK_CARD
        card.line.color.rgb = bcol
        card.line.width = Pt(1.5)

        tb = s4.shapes.add_textbox(bx + Inches(0.2), by + Inches(0.15), bw - Inches(0.4), bh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = btag
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = bcol

        p = tf.add_paragraph()
        p.space_before = Pt(4)
        p.text = bdesc
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_LIGHT

    # Bottom strip
    strip = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.6))
    strip.fill.solid()
    strip.fill.fore_color.rgb = COLOR_DARK_CARD
    strip.line.color.rgb = COLOR_DARK_LINE
    strip.line.width = Pt(1)

    tb_s = s4.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.333), Inches(0.4))
    tf_s = tb_s.text_frame
    p = tf_s.paragraphs[0]
    p.text = "AUTHORITY MODEL: Server-enforced limits · Models observe; pure code evaluates policy · Seller approval required for all actions"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL_LIGHT

    # ==========================================
    # SLIDE 5: Demo Scenarios Table
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, COLOR_BG_LIGHT)
    add_header(s5, "Demo Scenarios  ·  Benchmark Matrix",
               "Four return scenarios",
               "Expected prototype behavior across diverse real-world return situations.")
    add_footer(s5, "Consistent Demo Benchmark Suite · 4 Scenarios", 5)

    # Native PowerPoint Table
    rows = 5
    cols = 4
    table_shape = s5.shapes.add_table(rows, cols, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.5))
    table = table_shape.table

    table.columns[0].width = Inches(1.6)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(3.933)
    table.columns[3].width = Inches(4.0)

    headers = ["Scenario", "Case Identifier", "Evidence Situation", "Expected Prototype Behavior"]
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(236, 238, 233)
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.text = h_text
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEAL

    data_rows = [
        ("Ready", "Case A · ORD-1042\n₹1,499", "Seam tear verified at 00:14; inside 7-day window; brand tag matches.", "Prepare for seller approval → Generates mock label & refund draft in pending_return_receipt."),
        ("Mismatch", "Case B · ORD-8492\n₹1,499", "White ordered vs Black observed in frame 00:07; tag missing.", "Needs seller review → Flagged with proof pack for manual decision; no auto-denial, no label created."),
        ("Unclear", "Case C · ORD-5120\n₹850", "4-second dark video (brightness 28); only 2 usable frames.", "Request specific evidence → Prompts buyer for well-lit tag photo; preserves active case session."),
        ("Context", "Case D · ORD-3391\n₹2,199", "Continuous take tracking parcel from sealed label to unboxing.", "Check sequence continuity → Verifies chronological context; asks for missing angles if needed.")
    ]

    for r_idx, (scen, cid, ev, exp) in enumerate(data_rows, start=1):
        for c_idx, val in enumerate([scen, cid, ev, exp]):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = val
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_DARK
            if c_idx == 0:
                p.font.bold = True
                if scen == "Ready":
                    p.font.color.rgb = COLOR_TEAL
                elif scen == "Mismatch":
                    p.font.color.rgb = COLOR_AMBER

    # ==========================================
    # SLIDE 6: 24-Hour Delivery Plan
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, COLOR_BG_LIGHT)
    add_header(s6, "Delivery Plan  ·  Team Rookie",
               "A focused 24-hour prototype",
               "Realistic scope and role allocation for a 3-person sprint during the hackathon.")
    add_footer(s6, "Project-specific logic written inside the 24 hours · Checkable, buildable agent prototype", 6)

    # Left: 4 Phases
    lx = Inches(0.8)
    ly = Inches(1.9)
    lw = Inches(6.4)
    lh = Inches(4.8)

    card_tl = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, lx, ly, lw, lh)
    card_tl.fill.solid()
    card_tl.fill.fore_color.rgb = COLOR_WHITE
    card_tl.line.color.rgb = COLOR_LINE
    card_tl.line.width = Pt(1)

    tb_tl = s6.shapes.add_textbox(lx + Inches(0.25), ly + Inches(0.2), lw - Inches(0.5), lh - Inches(0.4))
    tf_tl = tb_tl.text_frame
    tf_tl.word_wrap = True
    tf_tl.margin_left = tf_tl.margin_top = tf_tl.margin_right = tf_tl.margin_bottom = 0

    p = tf_tl.paragraphs[0]
    p.text = "24-HOUR SPRINT TIMELINE"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    phases = [
        ("00h – 06h: Mock APIs & Core Flow", "FastAPI mock store and shipping services, SQLite seed data, core StateGraph."),
        ("06h – 12h: Evidence Pipeline & Policy Checks", "ffmpeg CPU frame extraction (~1 fps), multimodal prompt grounding, deterministic code policy rules."),
        ("12h – 18h: Integration & 4 Demo Scenarios", "End-to-end testing across Cases A, B, C, D; critic checks; fallback error handling."),
        ("18h – 24h: Demo Polish & Packaging", "Seller review console polish, 3-minute video walk-through, GitHub README, submission package.")
    ]
    for ptitle, pdesc in phases:
        p = tf_tl.add_paragraph()
        p.space_before = Pt(8)
        p.text = ptitle
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK

        p = tf_tl.add_paragraph()
        p.text = pdesc
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Team & Scope
    rx = Inches(7.5)
    ry = Inches(1.9)
    rw = Inches(5.0)
    rh = Inches(4.8)

    card_tm = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, ry, rw, rh)
    card_tm.fill.solid()
    card_tm.fill.fore_color.rgb = COLOR_WHITE
    card_tm.line.color.rgb = COLOR_LINE
    card_tm.line.width = Pt(1)

    tb_tm = s6.shapes.add_textbox(rx + Inches(0.25), ry + Inches(0.2), rw - Inches(0.5), rh - Inches(0.4))
    tf_tm = tb_tm.text_frame
    tf_tm.word_wrap = True
    tf_tm.margin_left = tf_tm.margin_top = tf_tm.margin_right = tf_tm.margin_bottom = 0

    p = tf_tm.paragraphs[0]
    p.text = "TEAM ROLES (3 MEMBERS)"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    p = tf_tm.add_paragraph()
    p.space_before = Pt(6)
    p.text = "• Person A (Backend & Logic): Agent flow, state graph, deterministic policy engine, critic, audit log.\n• Person B (Multimodal & UI): Video pipeline, frame indexer, seller console, evidence pack view.\n• Person C (QA & Integration): Integration, mock store data, 4 demo cases, fallback testing, demo video."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_DARK

    p = tf_tm.add_paragraph()
    p.space_before = Pt(14)
    p.text = "STRICT 24-HOUR SCOPE GUARDRAILS:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    p = tf_tm.add_paragraph()
    p.space_before = Pt(4)
    p.text = "• Mock services, staged evidence, single seller return workflow, no production payments.\n• Deliverable: Checkable, reproducible agent workflows with transparent action trail."
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED

    prs.save(output_pptx)
    print(f"Saved PPTX successfully: {output_pptx}")

if __name__ == "__main__":
    out_file = r"D:\Code\Code\AIO\Code\ZaloAIC\ReturnLens_v3.pptx"
    build_v3_deck(out_file)

